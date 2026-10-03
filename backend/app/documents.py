"""Extract plain text from uploaded request files (.txt, .md, .pdf, .docx)."""

from io import BytesIO
from zipfile import BadZipFile

from docx import Document
from pypdf import PdfReader
from pypdf.errors import PdfReadError

MAX_UPLOAD_BYTES = 10 * 1024 * 1024
MAX_TEXT_CHARS = 20_000
TEXT_ENCODINGS = ("utf-8-sig", "shift_jis", "cp1258", "latin-1")


class DocumentError(ValueError):
    """Raised with a Vietnamese, user-facing message."""


def decode_text(data: bytes) -> str:
    for encoding in TEXT_ENCODINGS:
        try:
            return data.decode(encoding)
        except UnicodeDecodeError:
            continue
    raise DocumentError("Không đọc được mã hóa ký tự của file.")


def _pdf_text(data: bytes) -> tuple[str, int]:
    try:
        reader = PdfReader(BytesIO(data))
        if reader.is_encrypted:
            raise DocumentError("PDF đang được đặt mật khẩu, hãy gỡ mật khẩu rồi tải lại.")
        pages = [page.extract_text() or "" for page in reader.pages]
    except DocumentError:
        raise
    except (PdfReadError, ValueError, KeyError) as exc:
        raise DocumentError("File PDF bị lỗi hoặc không đúng định dạng.") from exc
    return "\n\n".join(p.strip() for p in pages if p.strip()), len(pages)


def _docx_text(data: bytes) -> str:
    try:
        doc = Document(BytesIO(data))
    except (BadZipFile, KeyError, ValueError) as exc:
        raise DocumentError("File Word (.docx) bị lỗi hoặc không đúng định dạng.") from exc
    parts = [p.text for p in doc.paragraphs if p.text.strip()]
    for table in doc.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells]
            if any(cells):
                parts.append(" | ".join(cells))
    return "\n".join(parts)


def extract_text(filename: str, data: bytes) -> dict[str, object]:
    if len(data) > MAX_UPLOAD_BYTES:
        raise DocumentError("File quá lớn (tối đa 10 MB).")
    lower = filename.lower()
    pages = None
    if lower.endswith(".pdf"):
        text, pages = _pdf_text(data)
        if len(text.strip()) < 30:
            raise DocumentError(
                "Không trích xuất được chữ từ PDF. Có thể đây là bản scan (ảnh), chưa hỗ trợ OCR."
            )
    elif lower.endswith(".docx"):
        text = _docx_text(data)
    elif lower.endswith((".txt", ".md")):
        text = decode_text(data)
    else:
        raise DocumentError("Chỉ hỗ trợ file .pdf, .docx, .txt hoặc .md.")
    text = text.strip()
    truncated = len(text) > MAX_TEXT_CHARS
    return {
        "filename": filename,
        "text": text[:MAX_TEXT_CHARS],
        "pages": pages,
        "truncated": truncated,
    }
