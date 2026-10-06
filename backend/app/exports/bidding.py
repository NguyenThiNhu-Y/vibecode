"""Bidding workbook (docs/BIDDING_SPEC.md 6.1): Q&A, WBS, Summary, Master Schedule.

Sheet names and headers stay in English (sent as-is to Japanese / international clients); cell
content keeps the language of the data. Parent estimates, totals and the Summary matrix are Excel
formulas, so presales can edit a leaf estimate and every total follows.
"""

import io
import math
from datetime import date

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from app.agents.schedule import working_days_between
from app.exports.plan import monday, schedule_for, week_starts
from app.schemas.common import Phase, WorkType
from app.schemas.run import ScopingRun
from app.schemas.schedule import ScheduleResult
from app.schemas.wbs import PHASE_ORDER, WbsItem, WbsResult, children_map, ordered

PHASE = {"poc": "PoC", "mvp": "MVP", "production": "Production"}
PRIORITY = {"high": "High", "mid": "Mid", "low": "Low"}
PHASE_COLORS = {"poc": "F9A66C", "mvp": "F26F21", "production": "B5420E"}
PHASE_SOFT = {"poc": "FDE3D0", "mvp": "FBD0B5", "production": "EBC3B0"}
HEADER_FILL = PatternFill("solid", fgColor="D9D9D9")
HIGH_FILL = PatternFill("solid", fgColor="FDE2E1")
TOTAL_FILL = PatternFill("solid", fgColor="F2F2F2")
THIN = Side(style="thin", color="D9D9D9")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
WRAP = Alignment(wrap_text=True, vertical="top")
MAX_WIDTH = 60
WEEK_WIDTH = 6.5


# ---------- shared formatting ----------
def _header(ws: Worksheet, headers: list[str]) -> None:
    for col, title in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=col, value=title)
        cell.font, cell.fill, cell.border = Font(bold=True), HEADER_FILL, BORDER
        cell.alignment = Alignment(wrap_text=True, vertical="center")
    ws.freeze_panes = "A2"


def _finish(ws: Worksheet, columns: int, rows: int, fixed: dict[int, float] | None = None) -> None:
    """Auto-filter, borders, text wrap and column widths that fit the content (max 60)."""
    last = get_column_letter(columns)
    ws.auto_filter.ref = f"A1:{last}{max(rows, 1)}"
    widths = dict.fromkeys(range(1, columns + 1), 6.0)
    for row in ws.iter_rows(min_row=1, max_row=rows, max_col=columns):
        for cell in row:
            if cell.value is None:
                continue
            text = cell.value if isinstance(cell.value, str) else str(cell.value)
            size = 10 if text.startswith("=") else max(len(p) for p in text.split("\n"))
            widths[cell.column] = max(widths[cell.column], min(MAX_WIDTH, size + 2))
            cell.border = BORDER
            if cell.row > 1:
                cell.alignment = Alignment(
                    wrap_text=True, vertical="top", indent=cell.alignment.indent
                )
    widths.update(fixed or {})
    for col, width in widths.items():
        ws.column_dimensions[get_column_letter(col)].width = width


# ---------- Q&A ----------
def _qa(ws: Worksheet, run: ScopingRun) -> None:
    ws.title = "Q&A"
    headers = ["No", "Category", "Question", "Why we ask", "Priority", "Blocking",
               "Customer answer", "Status"]  # fmt: skip
    _header(ws, headers)
    questions = run.gaps.questions if run.gaps else []
    for n, q in enumerate(questions, start=1):
        answer = run.answers.get(q.id, "")
        values = [
            n,
            q.topic.value,
            q.question,
            q.why_it_matters,
            PRIORITY[q.priority.value],
            "Yes" if q.blocking else "No",
            answer,
            "Answered" if answer else "Open",
        ]
        for col, value in enumerate(values, start=1):
            cell = ws.cell(row=n + 1, column=col, value=value)
            if q.priority.value == "high":
                cell.fill = HIGH_FILL
    _finish(ws, len(headers), len(questions) + 1, {3: 55, 4: 45, 7: 45})


# ---------- WBS ----------
def _wbs_row(
    ws: Worksheet, r: int, item: WbsItem, kids: dict[str, list[WbsItem]], rows: dict[str, int]
) -> None:
    ws.cell(row=r, column=1, value=item.id)
    ws.cell(row=r, column=2, value=PHASE[item.phase.value])
    name = ws.cell(row=r, column=4 if item.level == 3 else 3, value=item.name)
    name.font = Font(bold=item.level == 1)
    if item.level == 2:
        name.alignment = Alignment(indent=1)
    if item.id in kids:  # parent: Excel formula over its direct children
        refs = ",".join(f"G{rows[c.id]}" for c in kids[item.id])
        ws.cell(row=r, column=7, value=f"=SUM({refs})").font = Font(bold=item.level == 1)
    else:  # only leaves carry Type, so Summary's SUMIFS never counts a parent twice
        ws.cell(row=r, column=5, value=item.type.value if item.type else None)
        ws.cell(row=r, column=6, value=PRIORITY[item.priority.value])
        ws.cell(row=r, column=7, value=item.estimate_md)
    ws.cell(row=r, column=8, value=", ".join(item.depends_on) or None)
    ws.cell(row=r, column=9, value=item.deliverable)
    note = "; ".join(filter(None, [item.note, ", ".join(t.value for t in item.tags)]))
    ws.cell(row=r, column=10, value=note or None)


def _wbs(ws: Worksheet, wbs: WbsResult) -> dict[str, int]:
    """Rows in phase order, a total row after each phase, a grand total. Returns the total row
    of every phase."""
    headers = ["ID", "Phase", "Task", "Sub-task", "Type", "Priority", "Estimate (MD)",
               "Depends on", "Deliverable", "Note"]  # fmt: skip
    _header(ws, headers)
    kids = children_map(wbs)
    plan: list[WbsItem | Phase] = []
    for phase in PHASE_ORDER:
        items = [i for i in ordered(wbs) if i.phase == phase]
        if items:
            plan += [*items, phase]
    rows = {e.id: r for r, e in enumerate(plan, start=2) if isinstance(e, WbsItem)}
    totals: dict[str, int] = {}
    for r, entry in enumerate(plan, start=2):
        if isinstance(entry, WbsItem):
            _wbs_row(ws, r, entry, kids, rows)
            continue
        label = PHASE[entry.value]
        groups = ",".join(f"G{rows[i.id]}" for i in wbs.items if i.phase == entry and i.level == 1)
        ws.cell(row=r, column=2, value=label)
        ws.cell(row=r, column=3, value=f"Total {label}").font = Font(bold=True)
        ws.cell(row=r, column=7, value=f"=SUM({groups})").font = Font(bold=True)
        for col in range(1, len(headers) + 1):
            ws.cell(row=r, column=col).fill = TOTAL_FILL
        totals[entry.value] = r
    r = len(plan) + 2
    ws.cell(row=r, column=3, value="Grand total").font = Font(bold=True)
    grand = ",".join(f"G{t}" for t in totals.values())
    ws.cell(row=r, column=7, value=f"=SUM({grand})").font = Font(bold=True)
    _finish(ws, len(headers), r, {3: 44, 4: 40, 9: 36, 10: 30})
    return totals


# ---------- Summary ----------
def _summary(ws: Worksheet, wbs: WbsResult) -> None:
    used = {leaf_type for t in wbs.totals for leaf_type in t.by_type}
    types = [t for t in WorkType if t in used]
    headers = ["Phase", *(t.value for t in types), "Total"]
    _header(ws, headers)
    phases = [t.phase for t in wbs.totals]
    last_type = get_column_letter(len(types) + 1)
    for r, phase in enumerate(phases, start=2):
        ws.cell(row=r, column=1, value=PHASE[phase.value]).font = Font(bold=True)
        for c in range(2, len(types) + 2):
            col = get_column_letter(c)
            ws.cell(row=r, column=c, value=f"=SUMIFS(WBS!$G:$G,WBS!$B:$B,$A{r},WBS!$E:$E,{col}$1)")
        ws.cell(row=r, column=len(types) + 2, value=f"=SUM(B{r}:{last_type}{r})").font = Font(
            bold=True
        )
    total_row = len(phases) + 2
    ws.cell(row=total_row, column=1, value="Total").font = Font(bold=True)
    for c in range(2, len(types) + 3):
        col = get_column_letter(c)
        cell = ws.cell(row=total_row, column=c, value=f"=SUM({col}2:{col}{total_row - 1})")
        cell.font, cell.fill = Font(bold=True), TOTAL_FILL
    ws.cell(row=total_row, column=1).fill = TOTAL_FILL
    _finish(ws, len(headers), total_row, {1: 14})
    for c in range(2, len(types) + 3):
        ws.column_dimensions[get_column_letter(c)].width = 10

    notes = [(f"Adjustment {PHASE[t.phase.value]}", [t.adjustment_note]) for t in wbs.totals
             if t.adjustment_note]  # fmt: skip
    r = total_row + 2
    span = max(len(headers), 8)
    for title, lines in [
        ("Assumptions", wbs.assumptions),
        ("Out of scope", wbs.out_of_scope),
        *notes,
    ]:
        if not lines:
            continue
        ws.cell(row=r, column=1, value=title).font = Font(bold=True)
        r += 1
        for line in lines:
            ws.cell(row=r, column=1, value=f"• {line}").alignment = WRAP
            ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=span)
            ws.row_dimensions[r].height = 15 * max(1, math.ceil(len(line) / 110))
            r += 1
        r += 1


# ---------- Master Schedule ----------
def _master(ws: Worksheet, wbs: WbsResult, schedule: ScheduleResult) -> None:
    headers = ["ID", "Task", "Phase", "Start", "End", "Days"]
    weeks = week_starts(schedule)
    _header(ws, headers + [w.strftime("%d/%m") for w in weeks])
    ws.freeze_panes = "G2"
    holidays = set(schedule.config.holidays)
    names = {i.id: i.name for i in wbs.items}
    by_phase = {i.id: i.phase.value for i in wbs.items}
    first_week = weeks[0]

    def bar(row: int, start: date, end: date, color: str) -> None:
        for k in range(
            (monday(start) - first_week).days // 7, (monday(end) - first_week).days // 7 + 1
        ):
            ws.cell(row=row, column=len(headers) + 1 + k).fill = PatternFill("solid", fgColor=color)

    r = 1
    for phase in schedule.phases:
        p = phase.phase.value
        r += 1
        values = ["", f"{PHASE[p]}", PHASE[p], phase.start, phase.end, phase.working_days]
        for col, value in enumerate(values, start=1):
            ws.cell(row=r, column=col, value=value).font = Font(bold=True)
        bar(r, phase.start, phase.end, PHASE_SOFT[p])
        for task in schedule.tasks:
            if by_phase[task.wbs_id] != p:
                continue
            r += 1
            days = working_days_between(task.start, task.end, holidays)
            for col, value in enumerate(
                [task.wbs_id, names[task.wbs_id], PHASE[p], task.start, task.end, days], start=1
            ):
                ws.cell(row=r, column=col, value=value)
            ws.cell(row=r, column=2).alignment = Alignment(indent=1)
            bar(r, task.start, task.end, PHASE_COLORS[p])
        for m in schedule.milestones:
            if m.phase != phase.phase:
                continue
            r += 1
            label = m.name + (f" ({m.payment_percent}%)" if m.payment_percent else "")
            for col, value in enumerate([m.id, label, PHASE[p], m.date, m.date, 0], start=1):
                ws.cell(row=r, column=col, value=value).font = Font(bold=True, color="B5420E")
            k = (monday(m.date) - first_week).days // 7
            mark = ws.cell(row=r, column=len(headers) + 1 + k, value="◆")
            mark.font, mark.alignment = (
                Font(bold=True, color="B5420E"),
                Alignment(horizontal="center"),
            )
    for row in ws.iter_rows(min_row=2, max_row=r, min_col=4, max_col=5):
        for cell in row:
            cell.number_format = "dd/mm/yyyy"
    _finish(ws, len(headers), r, {1: 7, 2: 42, 3: 11, 4: 11, 5: 11, 6: 6})
    for k in range(len(weeks)):
        col = len(headers) + 1 + k
        ws.column_dimensions[get_column_letter(col)].width = WEEK_WIDTH
        for row in range(1, r + 1):
            ws.cell(row=row, column=col).border = BORDER


def build_bidding_xlsx(run: ScopingRun) -> bytes:
    """Q&A only while there is no WBS yet (e.g. waiting for the client's answers)."""
    wb = Workbook()
    _qa(wb.active, run)
    if run.wbs:
        _wbs(wb.create_sheet("WBS"), run.wbs)
        _summary(wb.create_sheet("Summary"), run.wbs)
        schedule = schedule_for(run)
        if schedule:
            _master(wb.create_sheet("Master Schedule"), run.wbs, schedule)
    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()
