"""Word document: proposal Markdown converted to .docx + appendix tables (effort, WBS, matrix)."""

import io
import re

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor

from app.company import match_case_studies, placeholders, render_blocks
from app.exports.common import (
    ACCENT,
    COVERAGE_LABELS,
    PHASE_LABELS,
    requirement_texts,
)
from app.exports.i18n import LABELS, format_money
from app.exports.templating import fill_docx, find_body_marker, move_after_marker
from app.schemas.run import ScopingRun
from app.templates_store import ExportKit

_BOLD = re.compile(r"(\*\*[^*]+\*\*)")


def _fonts(doc: Document) -> None:
    style = doc.styles["Normal"]
    style.font.name = "Arial"
    style.font.size = Pt(10.5)
    rpr = style.element.get_or_add_rPr()
    fonts = rpr.find(qn("w:rFonts"))
    if fonts is None:
        fonts = OxmlElement("w:rFonts")
        rpr.append(fonts)
    fonts.set(qn("w:eastAsia"), "Yu Gothic")
    for level in (1, 2, 3):
        doc.styles[f"Heading {level}"].font.color.rgb = RGBColor.from_string(
            ACCENT if level == 1 else "0C111D"
        )


def _inline(paragraph, text: str) -> None:
    for part in _BOLD.split(text):
        if part.startswith("**") and part.endswith("**"):
            paragraph.add_run(part[2:-2]).bold = True
        elif part:
            paragraph.add_run(part.replace("`", ""))


def _shade(cell, hex_color: str) -> None:
    shading = OxmlElement("w:shd")
    shading.set(qn("w:val"), "clear")
    shading.set(qn("w:fill"), hex_color)
    cell._tc.get_or_add_tcPr().append(shading)


def _has_style(doc: Document, name: str) -> bool:
    return any(s.name == name for s in doc.styles)


def _heading(doc: Document, text: str, level: int) -> None:
    """Heading style when the (company) template has it, bold paragraph otherwise."""
    if _has_style(doc, f"Heading {level}"):
        doc.add_heading(text, level=level)
    else:
        run = doc.add_paragraph().add_run(text)
        run.bold, run.font.size = True, Pt({1: 16, 2: 13}.get(level, 11.5))


def _para(doc: Document, style: str | None = None):
    return doc.add_paragraph(style=style if style and _has_style(doc, style) else None)


def _accent(doc: Document) -> str:
    """Table header color: the template's Heading 1 color when it sets one."""
    try:
        color = doc.styles["Heading 1"].font.color.rgb
    except (KeyError, AttributeError):
        color = None
    return str(color) if color is not None and str(color) != "000000" else ACCENT


def _table(doc: Document, header: list[str], rows: list[list[str]]) -> None:
    table = doc.add_table(rows=1, cols=len(header))
    if _has_style(doc, "Table Grid"):
        table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    accent = _accent(doc)
    for cell, text in zip(table.rows[0].cells, header, strict=True):
        cell.text = ""
        run = cell.paragraphs[0].add_run(text)
        run.bold = True
        run.font.color.rgb = RGBColor(255, 255, 255)
        _shade(cell, accent)
    for values in rows:
        cells = table.add_row().cells
        for cell, value in zip(cells, values, strict=True):
            cell.text = str(value)
    doc.add_paragraph()


def _markdown(doc: Document, markdown: str) -> None:
    table_lines: list[str] = []

    def flush_table() -> None:
        if not table_lines:
            return
        rows = [
            [c.strip() for c in line.strip().strip("|").split("|")]
            for line in table_lines
            if not re.match(r"^\s*\|?\s*:?-{2,}", line)
        ]
        table_lines.clear()
        if rows:
            width = len(rows[0])
            _table(doc, rows[0], [(r + [""] * width)[:width] for r in rows[1:]])

    for line in markdown.splitlines():
        stripped = line.strip()
        if stripped.startswith("|"):
            table_lines.append(stripped)
            continue
        flush_table()
        if not stripped:
            continue
        heading = re.match(r"^(#{1,6})\s+(.*)", stripped)
        if heading:
            _heading(doc, heading.group(2).replace("**", ""), min(len(heading.group(1)), 3))
        elif re.match(r"^[-*]\s+", stripped):
            _inline(_para(doc, "List Bullet"), re.sub(r"^[-*]\s+", "", stripped))
        elif re.match(r"^\d+\.\s+", stripped):
            _inline(_para(doc, "List Number"), re.sub(r"^\d+\.\s+", "", stripped))
        else:
            _inline(doc.add_paragraph(), stripped)
    flush_table()


def _blocks(doc: Document, run: ScopingRun, kit: ExportKit, position: str, lang: str) -> None:
    for title, body in render_blocks(run, kit.company, kit.blocks, position, lang):
        _heading(doc, title, 1 if position == "start" else 2)
        for line in body.splitlines():
            if line.strip():
                _inline(_para(doc), line.strip())


def _case_studies(doc: Document, run: ScopingRun, kit: ExportKit, lang: str) -> None:
    matches = match_case_studies(run, kit.case_studies)
    if not matches:
        return
    L = LABELS[lang]
    _heading(doc, L["case_studies"], 2)
    note = _para(doc)
    note.add_run(L["cs_sub"]).italic = True
    for match in matches:
        cs = match["case"]
        _heading(doc, cs.title, 3)
        meta = " · ".join(
            str(v) for v in (cs.industry, L["markets"][cs.market], cs.year, cs.duration) if v
        )
        _para(doc).add_run(meta).italic = True
        for label, text in ((L["cs_challenge"], cs.challenge), (L["cs_solution"], cs.solution)):
            p = _para(doc)
            p.add_run(f"{label}: ").bold = True
            p.add_run(text)
        p = _para(doc)
        p.add_run(f"{L['cs_results']}: ").bold = True
        p.add_run("; ".join(cs.results))


def _quotation(doc: Document, run: ScopingRun) -> None:
    q = run.quotation
    if q is None:
        return
    money = lambda v: format_money(v, q.currency)  # noqa: E731
    _heading(doc, f"E. Báo giá sơ bộ – {LABELS['vi']['contract'][q.contract_model]}", 2)
    if q.contract_model == "odc":
        rows = [[m.role_label, f"{m.fte:g}", money(m.monthly_cost)] for m in q.odc_team]
        rows.append(["Chi phí đội / tháng", "", money(q.odc_monthly_cost or 0)])
        rows.append([f"Tổng {q.months} tháng", "", money(q.total)])
        _table(doc, ["Vai trò", "FTE", "Chi phí / tháng"], rows)
    else:
        rows = [
            [
                PHASE_LABELS[p.phase.value],
                str(p.person_days),
                money(p.amount),
                f"{money(p.min_amount)} – {money(p.max_amount)}",
            ]
            for p in q.phases
        ]
        if q.contract_model == "fixed_price":
            rows.append([f"Dự phòng rủi ro {q.contingency_pct:g}%", "", money(q.contingency), ""])
        label = (
            "Tổng (đã gồm dự phòng)" if q.contract_model == "fixed_price" else "Tổng dự toán (T&M)"
        )
        rows.append([label, "", money(q.total), f"{money(q.total_min)} – {money(q.total_max)}"])
        _table(doc, ["Giai đoạn", "Ngày công", "Chi phí", "Khoảng ước tính"], rows)
    _table(
        doc,
        ["Mốc thanh toán", "%", "Số tiền"],
        [[m.name, f"{m.percent:g}%", money(m.amount)] for m in q.milestones],
    )
    for a in q.assumptions:
        _inline(_para(doc, "List Bullet"), a)


def build_docx(run: ScopingRun, kit: ExportKit | None = None) -> bytes:
    """Builtin layout, or the company Word template (placeholders filled, body at the marker)."""
    kit = kit or ExportKit.load()
    template = kit.templates.get("proposal_docx")
    lang = run.proposal.language if run.proposal else "vi"
    values = placeholders(run, kit.company, lang)
    if template:
        doc = Document(str(template))
        fill_docx(doc, values)
        marker = find_body_marker(doc)
        first_new = len(doc.element.body) - 1  # new content is appended before the final sectPr
    else:
        doc = Document()
        _fonts(doc)
        marker = None
        footer = doc.sections[0].footer.paragraphs[0]
        footer.text = f"{kit.company.name} · {kit.company.confidential_footer}".strip(" ·")
        footer.runs[0].font.size = Pt(8)

    _blocks(doc, run, kit, "start", lang)
    if run.proposal:
        _markdown(doc, run.proposal.markdown)
    _case_studies(doc, run, kit, lang)
    _blocks(doc, run, kit, "end", lang)

    doc.add_page_break()
    _heading(doc, "Phụ lục", 1)
    note = _para(doc)
    note.add_run(
        "Các bảng dưới đây được ScopeAI tính bằng code từ kết quả phân tích. "
    ).italic = True

    if run.architecture:
        _heading(doc, "A. Effort theo giai đoạn", 2)
        basis = run.effort_basis
        if basis:
            factors = " × ".join(f"{m.factor} ({m.label})" for m in basis.multipliers) or "1"
            _para(doc).add_run(f"Công thức: bảng chuẩn × {factors} = × {basis.factor}")
        _table(
            doc,
            ["Giai đoạn", "Ngày công", "Đội ngũ", "Bàn giao"],
            [
                [
                    PHASE_LABELS[e.phase.value],
                    f"{e.min_person_days}–{e.max_person_days}",
                    ", ".join(e.team),
                    "; ".join(e.deliverables),
                ]
                for e in run.architecture.estimates
            ],
        )
        _heading(doc, "B. Thành phần kiến trúc", 2)
        _table(
            doc,
            ["Thành phần", "Mục đích", "Công nghệ"],
            [[c.name, c.purpose, ", ".join(c.tech_options)] for c in run.architecture.components],
        )

    if run.wbs:
        _heading(doc, "C. WBS", 2)
        sched = {t.id: t for t in run.schedule.tasks} if run.schedule else {}
        _table(
            doc,
            ["ID", "Giai đoạn", "Đầu việc", "Vai trò", "Ngày công", "Tuần"],
            [
                [
                    t.id,
                    PHASE_LABELS[t.phase.value],
                    t.name,
                    t.role,
                    str(t.person_days),
                    f"T{sched[t.id].start_day // 5 + 1}–T{(sched[t.id].end_day - 1) // 5 + 1}"
                    if t.id in sched
                    else "",
                ]
                for t in run.wbs.tasks
            ],
        )

    if run.requirements and not run.requirements.skipped:
        _heading(doc, "D. Bảng đáp ứng yêu cầu", 2)
        texts = requirement_texts(run)
        _table(
            doc,
            ["Mã", "Yêu cầu", "Mức đáp ứng", "Ghi chú"],
            [
                [i.req_id, texts.get(i.req_id, ("",))[0], COVERAGE_LABELS[i.coverage], i.note]
                for i in run.requirements.items
            ],
        )

    _quotation(doc, run)

    if template and marker is not None:
        move_after_marker(doc, marker, first_new)
    buffer = io.BytesIO()
    doc.save(buffer)
    return buffer.getvalue()
