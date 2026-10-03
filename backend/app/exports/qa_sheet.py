"""Q&A sheet (質問票) for the customer, in the customer's language, and its re-import."""

import io
from datetime import date

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation

from app.company import placeholders
from app.documents import DocumentError
from app.exports.common import ACCENT
from app.exports.templating import fill_workbook
from app.schemas.run import ScopingRun
from app.schemas.settings import SheetMapping
from app.templates_store import ExportKit

LABELS = {
    "vi": {
        "title": "DANH SÁCH CÂU HỎI LÀM RÕ",
        "headers": ["ID", "STT", "Chủ đề", "Câu hỏi", "Lý do cần hỏi", "Bắt buộc", "Câu trả lời của Quý khách"],
        "yes": "Có",
        "note": "Vui lòng điền câu trả lời vào cột cuối (màu vàng) và gửi lại file này. Không sửa cột ID.",
        "topics": {
            "data": "Dữ liệu", "users": "Người dùng", "accuracy": "Độ chính xác", "infra": "Hạ tầng",
            "budget": "Ngân sách", "timeline": "Thời gian", "compliance": "Tuân thủ", "integration": "Tích hợp",
        },
    },
    "en": {
        "title": "CLARIFICATION QUESTIONS",
        "headers": ["ID", "No.", "Topic", "Question", "Why we ask", "Required", "Your answer"],
        "yes": "Yes",
        "note": "Please fill in the last column (yellow) and send this file back. Do not edit the ID column.",
        "topics": {
            "data": "Data", "users": "Users", "accuracy": "Accuracy", "infra": "Infrastructure",
            "budget": "Budget", "timeline": "Timeline", "compliance": "Compliance", "integration": "Integration",
        },
    },
    "ja": {
        "title": "質問票",
        "headers": ["ID", "No.", "分類", "ご質問", "質問の背景", "必須", "ご回答"],
        "yes": "必須",
        "note": "最終列（黄色）にご回答をご記入のうえ、本ファイルをご返送ください。ID列は変更しないでください。",
        "topics": {
            "data": "データ", "users": "利用者", "accuracy": "精度", "infra": "インフラ",
            "budget": "予算", "timeline": "スケジュール", "compliance": "コンプライアンス", "integration": "システム連携",
        },
    },
}  # fmt: skip
ANSWER_HEADERS = {labels["headers"][-1] for labels in LABELS.values()}
HEADER_ROW = 4


def _template_sheet(run: ScopingRun, kit: ExportKit, template) -> bytes:
    """Company Q&A template: placeholders filled, questions written by the column mapping."""
    lang = run.intake.language if run.intake else "vi"
    labels = LABELS[lang]
    mapping = kit.config.qa_sheet
    wb = load_workbook(str(template))
    fill_workbook(wb, placeholders(run, kit.company, lang))
    ws = wb[mapping.sheet] if mapping.sheet in wb.sheetnames else wb.create_sheet(mapping.sheet)
    yellow = PatternFill("solid", fgColor="FFF7CC")
    for n, q in enumerate(run.gaps.questions, start=1):  # type: ignore[union-attr]
        row = mapping.start_row + n - 1
        values = {
            "id": q.id, "num": n, "topic": labels["topics"][q.topic.value], "question": q.question,
            "why": q.why_it_matters, "blocking": labels["yes"] if q.blocking else "",
            "answer": run.answers.get(q.id, ""),
        }  # fmt: skip
        for key, col in mapping.columns.items():
            if key in values:
                cell = ws[f"{col}{row}"]
                cell.value = values[key]
                cell.alignment = Alignment(wrap_text=True, vertical="top")
        if "answer" in mapping.columns:
            ws[f"{mapping.columns['answer']}{row}"].fill = yellow
    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def build_qa_sheet(run: ScopingRun, kit: ExportKit | None = None) -> bytes:
    if run.gaps is None:
        raise DocumentError("Run chưa có câu hỏi làm rõ.")
    kit = kit or ExportKit.load()
    if kit.templates.get("qa_sheet"):
        return _template_sheet(run, kit, kit.templates["qa_sheet"])
    lang = run.intake.language if run.intake else "vi"
    labels = LABELS[lang]
    wb = Workbook()
    ws = wb.active
    ws.title = "Q&A"
    ws["A1"] = (
        f"{labels['title']} – {run.project_name or (run.intake.business_goal[:60] if run.intake else '')}"
    )
    ws["A1"].font = Font(bold=True, size=14, color=ACCENT)
    ws["A2"] = f"{date.today():%Y-%m-%d} · Ref: {run.id}"
    ws["A3"] = labels["note"]
    ws["A3"].font = Font(italic=True, color="6B7280")
    thin = Side(style="thin", color="D9D9D9")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    headers = (
        labels["headers"]
        if lang == "vi"
        else [h for i, h in enumerate(labels["headers"]) if i != 4]
    )
    widths = [7, 6, 16, 60, 40, 10, 50] if lang == "vi" else [7, 6, 16, 60, 10, 50]
    for col, (title, width) in enumerate(zip(headers, widths, strict=True), start=1):
        cell = ws.cell(row=HEADER_ROW, column=col, value=title)
        cell.fill = PatternFill("solid", fgColor=ACCENT)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.border = border
        ws.column_dimensions[cell.column_letter].width = width
    yellow = PatternFill("solid", fgColor="FFF7CC")
    for n, q in enumerate(run.gaps.questions, start=1):
        values = [q.id, n, labels["topics"][q.topic.value], q.question]
        if lang == "vi":
            values.append(q.why_it_matters)
        values += [labels["yes"] if q.blocking else "", run.answers.get(q.id, "")]
        row = HEADER_ROW + n
        for col, value in enumerate(values, start=1):
            cell = ws.cell(row=row, column=col, value=value)
            cell.border = border
            cell.alignment = Alignment(wrap_text=True, vertical="top")
        ws.cell(row=row, column=len(values)).fill = yellow
        ws.row_dimensions[row].height = 45
    ws.freeze_panes = ws.cell(row=HEADER_ROW + 1, column=1)
    ws.protection.sheet = False
    ws.add_data_validation(
        DataValidation(type="textLength", operator="lessThanOrEqual", formula1="2000")
    )
    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def _by_mapping(wb, mapping: SheetMapping) -> dict[str, str] | None:
    if mapping.sheet not in wb.sheetnames or not {"id", "answer"} <= set(mapping.columns):
        return None
    ws = wb[mapping.sheet]
    answers = {}
    for row in range(mapping.start_row, ws.max_row + 1):
        qid = ws[f"{mapping.columns['id']}{row}"].value
        answer = ws[f"{mapping.columns['answer']}{row}"].value
        if qid is not None and answer is not None and str(answer).strip():
            answers[str(qid).strip()] = str(answer).strip()
    return answers


def parse_qa_answers(data: bytes, mapping: SheetMapping | None = None) -> dict[str, str]:
    """Read answers back from a filled Q&A sheet: ScopeAI layout (any of the three languages)
    or, failing that, the company template's column mapping."""
    try:
        wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except Exception as exc:
        raise DocumentError("File Q&A không đúng định dạng .xlsx.") from exc
    if mapping is not None and mapping.sheet in wb.sheetnames:
        ws = wb[mapping.sheet]
    else:
        ws = wb.worksheets[0]
    rows = [list(r) for r in ws.iter_rows(values_only=True)]
    for i, row in enumerate(rows[:15]):
        headers = [str(c).strip() if c is not None else "" for c in row]
        if headers and headers[0] == "ID" and any(h in ANSWER_HEADERS for h in headers):
            answer_col = next(j for j, h in enumerate(headers) if h in ANSWER_HEADERS)
            answers = {}
            for data_row in rows[i + 1 :]:
                qid = str(data_row[0]).strip() if data_row and data_row[0] is not None else ""
                answer = data_row[answer_col] if answer_col < len(data_row) else None
                if qid and answer is not None and str(answer).strip():
                    answers[qid] = str(answer).strip()
            return answers
    if mapping is not None:
        wb.close()
        wb = load_workbook(io.BytesIO(data), data_only=True)
        found = _by_mapping(wb, mapping)
        if found is not None:
            return found
    raise DocumentError("Không tìm thấy bảng câu hỏi (cột ID và cột câu trả lời) trong file.")
