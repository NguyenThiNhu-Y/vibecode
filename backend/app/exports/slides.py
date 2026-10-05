"""Proposal deck (.pptx) in vi/en/ja: editable shapes for the architecture diagram and Gantt.

Fixed labels come from app/exports/i18n.py; agent-written content goes through `tr`, a mapping
filled by app/exports/translation.py (identity for Vietnamese)."""

import io
from collections.abc import Callable
from datetime import datetime

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE, PP_PLACEHOLDER
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.xmlchemy import OxmlElement
from pptx.slide import Slide
from pptx.util import Emu, Inches, Pt

from app.company import load_case_studies, match_case_studies, placeholders, render_blocks
from app.exports.common import (
    ACCENT,
    COVERAGE_COLORS,
    INK,
    PHASE_LABELS,
    coverage_counts,
    graph_for,
    layers,
    requirement_texts,
)
from app.exports.i18n import LABELS, format_money
from app.exports.templating import fit_logo, mix, open_pptx_template, tidy_placeholders
from app.schemas.run import ScopingRun
from app.templates_store import ExportKit

W, H = Inches(13.333), Inches(7.5)
FONTS = {"vi": "Arial", "en": "Arial", "ja": "Yu Gothic"}
GRAY = "6B7280"
LIGHT = "F5F5F4"
# Body area of the design space (everything below the title), mapped onto a template's content area
DESIGN_LEFT, DESIGN_TOP = Inches(0.75), Inches(0.9)
DESIGN_W, DESIGN_H = Inches(11.85), Inches(6.1)


def _rgb(hex_color: str) -> RGBColor:
    return RGBColor.from_string(hex_color)


def _cut(text: str | None, limit: int) -> str:
    text = (text or "").strip()
    return text if len(text) <= limit else text[: limit - 1].rstrip() + "…"


def deck_texts(run: ScopingRun) -> list[str]:
    """Agent-written strings shown on slides (what needs translating for non-vi decks)."""
    texts: list[str] = []
    if run.intake:
        i = run.intake
        texts += [i.business_goal, i.current_process or "", i.budget or "", i.timeline or ""]
        texts += i.users + i.data_sources + i.constraints
    if run.pattern:
        texts += [run.pattern.rationale] + [r.reason for r in run.pattern.rejected]
        texts += run.pattern.assumptions
    if run.feasibility:
        for r in run.feasibility.risks:
            texts += [r.description, r.mitigation]
    if run.architecture:
        texts += list(graph_for(run).nodes.values())
        for e in run.architecture.estimates:
            texts += e.team + e.deliverables
        texts += run.architecture.out_of_scope
    if run.wbs:
        texts += [t.name for t in run.wbs.tasks] + [t.role for t in run.wbs.tasks]
    if run.requirements:
        texts += [i.note for i in run.requirements.items]
    if run.quotation:
        texts += run.quotation.assumptions + [m.name for m in run.quotation.milestones]
    for att in run.attachments:
        for p in att.data_profiles:
            texts += p.issues
    if run.quotation:
        texts += [m.role_label for m in run.quotation.odc_team]
    for match in match_case_studies(run, load_case_studies()):
        cs = match["case"]
        texts += [cs.title, cs.industry, cs.challenge, cs.solution, *cs.results]
    return list(dict.fromkeys(t for t in texts if t and t.strip()))


class Deck:
    """Draws every slide in a fixed 13.33 x 7.5 in design space. With a company template the
    design-space body area is mapped (uniform scale + offset) onto the template's content area,
    titles go into the template's title placeholder and colors follow its theme."""

    def __init__(
        self,
        run: ScopingRun,
        lang: str = "vi",
        tr: Callable[[str], str] | None = None,
        kit: ExportKit | None = None,
    ):
        self.run = run
        self.lang = lang if lang in LABELS else "vi"
        self.L = LABELS[self.lang]
        self.t = tr or (lambda s: s)
        self.font = FONTS[self.lang]
        self.kit = kit or ExportKit.load()
        self.values = placeholders(run, self.kit.company, self.lang)
        self.logo = self.kit.logo
        self.count = 0
        self.tpl = None
        path = self.kit.templates.get("slides")
        if path:
            self.tpl = open_pptx_template(path, self.kit.config.pptx, self.values, self.logo)
            self.prs = self.tpl.prs
            left, top, width, height = self.tpl.region
            self.k = min(width / DESIGN_W, height / DESIGN_H)
            self.ox, self.oy = left + (width - DESIGN_W * self.k) / 2, top
        else:
            self.prs = Presentation()
            self.prs.slide_width, self.prs.slide_height = W, H
            self.blank = self.prs.slide_layouts[6]
            self.k, self.ox, self.oy = 1.0, DESIGN_LEFT, DESIGN_TOP
        accent = (self.tpl.accent if self.tpl else None) or ACCENT
        self.accent = accent
        self.soft, self.soft2 = mix(accent, "FFFFFF", 0.9), mix(accent, "FFFFFF", 0.94)
        self.strong = mix(accent, "000000", 0.3)
        self.phase_colors = {
            "poc": mix(accent, "FFFFFF", 0.45),
            "mvp": accent,
            "production": self.strong,
        }

    # ---------- design space -> slide ----------
    def X(self, v) -> int:
        return int(self.ox + (v - DESIGN_LEFT) * self.k)

    def Y(self, v) -> int:
        return int(self.oy + (v - DESIGN_TOP) * self.k)

    def S(self, v) -> int:
        return int(v * self.k)

    def pt(self, size: float) -> Pt:
        return Pt(round(size * self.k, 1))

    # ---------- primitives ----------
    def text(
        self,
        slide: Slide,
        x,
        y,
        w,
        h,
        lines,
        size=14,
        color=INK,
        bold=False,
        align=None,
        bullets=False,
        space=6,
    ):
        box = slide.shapes.add_textbox(self.X(x), self.Y(y), self.S(w), self.S(h))
        frame = box.text_frame
        frame.word_wrap = True
        for i, line in enumerate([lines] if isinstance(lines, str) else lines):
            p = frame.paragraphs[0] if i == 0 else frame.add_paragraph()
            run = p.add_run()
            run.text = f"•  {line}" if bullets else line
            run.font.size, run.font.bold, run.font.name = self.pt(size), bold, self.font
            run.font.color.rgb = _rgb(color)
            if align:
                p.alignment = align
            if bullets or space != 6:
                p.space_after = self.pt(space)
        return box

    def rect(self, slide: Slide, x, y, w, h, fill, line=None, shape=MSO_SHAPE.RECTANGLE):
        s = slide.shapes.add_shape(shape, self.X(x), self.Y(y), self.S(w), self.S(h))
        s.fill.solid()
        s.fill.fore_color.rgb = _rgb(fill)
        if line:
            s.line.color.rgb = _rgb(line)
            s.line.width = Pt(1.25)
        else:
            s.line.fill.background()
        s.shadow.inherit = False
        return s

    def rounded(self, slide: Slide, x, y, w, h, fill, radius=0.15):
        shape = self.rect(slide, x, y, w, h, fill, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
        shape.adjustments[0] = radius
        return shape

    def slide(self, title: str, subtitle: str | None = None) -> Slide:
        self.count += 1
        if self.tpl:
            s = self.prs.slides.add_slide(self.tpl.content_layout)
            if s.shapes.title is not None:
                s.shapes.title.text = title
            else:
                self.text(
                    s,
                    Inches(0.74),
                    Inches(0.9),
                    Inches(11.5),
                    Inches(0.5),
                    title,
                    size=24,
                    bold=True,
                )
            tidy_placeholders(s, self.tpl.content_layout)
            if subtitle:
                self.text(
                    s,
                    Inches(0.74),
                    Inches(0.9),
                    Inches(11.5),
                    Inches(0.4),
                    subtitle,
                    size=13,
                    color=GRAY,
                )
            return s
        s = self.prs.slides.add_slide(self.blank)
        self.rect(s, Inches(0.5), Inches(0.45), Inches(0.09), Inches(0.62), self.accent)
        title_w = 11.5
        if self.logo:  # top-right corner; the title box makes room for it
            title_w = 10.6
            box = fit_logo(
                Inches(11.55), Inches(0.3), Inches(1.3), Inches(0.52), self.logo, 0, "right"
            )
            s.shapes.add_picture(str(self.logo), *box)
        self.text(
            s, Inches(0.72), Inches(0.33), Inches(title_w), Inches(0.6), title, size=26, bold=True
        )
        if subtitle:
            self.text(
                s,
                Inches(0.74),
                Inches(0.9),
                Inches(11.5),
                Inches(0.4),
                subtitle,
                size=13,
                color=GRAY,
            )
        footer = _cut(self.run.proposal.title if self.run.proposal else "ScopeAI", 90)
        self.text(s, Inches(0.5), Inches(7.0), Inches(10), Inches(0.3), footer, size=9, color=GRAY)
        self.text(
            s,
            Inches(12.2),
            Inches(7.0),
            Inches(0.7),
            Inches(0.3),
            str(self.count),
            size=9,
            color=GRAY,
            align=PP_ALIGN.RIGHT,
        )
        return s

    def table(self, slide: Slide, x, y, w, header, rows, widths, size=11, bold_last=False):
        shape = slide.shapes.add_table(
            len(rows) + 1,
            len(header),
            self.X(x),
            self.Y(y),
            self.S(w),
            self.S(Inches(0.4) * (len(rows) + 1)),
        )
        table = shape.table
        for i, width in enumerate(widths):
            table.columns[i].width = self.S(Inches(width))
        for r, values in enumerate([header] + rows):
            last = bold_last and r == len(rows)
            for c, value in enumerate(values):
                cell = table.cell(r, c)
                cell.text = str(value)
                cell.margin_top = cell.margin_bottom = Inches(0.04)
                for p in cell.text_frame.paragraphs:
                    for run in p.runs:
                        run.font.size, run.font.name = self.pt(size), self.font
                        run.font.bold = r == 0 or last
                        run.font.color.rgb = _rgb("FFFFFF" if r == 0 else INK)
                cell.fill.solid()
                cell.fill.fore_color.rgb = _rgb(
                    self.accent if r == 0 else self.soft if last else ("FFFFFF" if r % 2 else LIGHT)
                )
        return table

    # ---------- slides ----------
    def template_cover(self) -> None:
        run, L = self.run, self.L
        layout = self.tpl.cover_layout
        s = self.prs.slides.add_slide(layout)
        title = _cut(run.proposal.title if run.proposal else "Proposal", 110)
        status = L["approved"] if run.status.value == "approved" else L["draft"]
        meta = " · ".join(x for x in (run.client_name, f"{datetime.now():%Y-%m-%d}", status) if x)
        goal = _cut(self.t(run.intake.business_goal), 220) if run.intake else ""
        subtitle = next(
            (p for p in s.placeholders if p.placeholder_format.type == PP_PLACEHOLDER.SUBTITLE),
            None,
        )
        if s.shapes.title is not None:
            s.shapes.title.text = title
        if subtitle is not None:
            subtitle.text_frame.text = goal
            subtitle.text_frame.add_paragraph().text = meta
        tidy_placeholders(s, layout)
        if s.shapes.title is None or subtitle is None:  # layout without placeholders: draw text
            size = self.prs.slide_width
            box = s.shapes.add_textbox(
                int(size * 0.08), int(self.prs.slide_height * 0.35), int(size * 0.84), Inches(2)
            )
            box.text_frame.word_wrap = True
            box.text_frame.text = title if s.shapes.title is None else ""
            box.text_frame.add_paragraph().text = f"{goal}\n{meta}"

    def title_slide(self) -> None:
        self.count += 1
        if self.tpl:
            self.template_cover()
            return
        s = self.prs.slides.add_slide(self.blank)
        run, L = self.run, self.L
        self.rect(s, 0, 0, W, H, "0B0F1C")
        if self.logo:  # on a white card so dark logos stay visible on the dark cover
            x, y, cx, cy = fit_logo(
                Inches(1.0), Inches(0.7), Inches(2.6), Inches(0.6), self.logo, 0, "left"
            )
            pad = Inches(0.16)
            self.rounded(s, x - pad, y - pad, cx + 2 * pad, cy + 2 * pad, "FFFFFF", 0.12)
            s.shapes.add_picture(str(self.logo), x, y, cx, cy)
        title = run.proposal.title if run.proposal else "Proposal"
        # Estimate wrapped lines (~34 chars/line at 38pt, ~42 at 32pt) to place the goal below.
        size = 38 if len(title) <= 34 else 32
        lines = min(3, -(-len(title) // (34 if size == 38 else 42)))
        title_h = lines * size * 1.3 / 72 + 0.15
        title_y = 2.1 if lines == 1 else 1.7
        goal_y = title_y + title_h + 0.3
        self.rect(
            s,
            Inches(0.8),
            Inches(title_y + 0.1),
            Inches(0.14),
            Inches(goal_y - title_y + 0.6),
            self.accent,
        )
        self.text(
            s,
            Inches(1.15),
            Inches(title_y),
            Inches(11),
            Inches(title_h),
            _cut(title, 110),
            size=size,
            color="FFFFFF",
            bold=True,
        )
        if run.intake:
            self.text(
                s,
                Inches(1.17),
                Inches(goal_y),
                Inches(10.5),
                Inches(1.0),
                _cut(self.t(run.intake.business_goal), 220),
                size=17,
                color="C9CED9",
            )
        status = L["approved"] if run.status.value == "approved" else L["draft"]
        who = f"{run.client_name} · " if run.client_name else ""
        self.text(
            s,
            Inches(1.17),
            Inches(5.6),
            Inches(10),
            Inches(0.4),
            f"{who}{datetime.now():%Y-%m-%d} · {status}",
            size=13,
            color="FFB27F",
        )
        self.text(
            s,
            Inches(1.17),
            Inches(6.6),
            Inches(6),
            Inches(0.4),
            L["made_by"],
            size=11,
            color="6C7591",
        )

    def context_slide(self) -> None:
        intake, L, t = self.run.intake, self.L, self.t
        if not intake:
            return
        s = self.slide(L["context"])
        self.text(
            s,
            Inches(0.75),
            Inches(1.35),
            Inches(11.8),
            Inches(1.0),
            _cut(t(intake.business_goal), 260),
            size=20,
            bold=True,
        )
        unknown = L["unknown"]
        values = [
            t(intake.current_process) if intake.current_process else unknown,
            ", ".join(t(u) for u in intake.users) or unknown,
            ", ".join(t(d) for d in intake.data_sources) or unknown,
            ", ".join(t(c) for c in intake.constraints) or unknown,
            f"{t(intake.budget) if intake.budget else unknown} / {t(intake.timeline) if intake.timeline else unknown}",
        ]
        rows = [[k, _cut(v, 160)] for k, v in zip(L["fields"], values, strict=True)]
        self.table(
            s,
            Inches(0.75),
            Inches(2.6),
            Inches(11.8),
            [L["item"], L["info"]],
            rows,
            [3.0, 8.8],
            size=13,
        )

    def inputs_slide(self) -> None:
        L, t = self.L, self.t
        data = [(a, p) for a in self.run.attachments for p in a.data_profiles]
        code = [a for a in self.run.attachments if a.code_profile]
        reqs = sum(len(a.requirements) for a in self.run.attachments)
        docs = [a for a in self.run.attachments if a.kind == "document"]
        if not (data or code or reqs or docs):
            return
        s = self.slide(L["inputs"], L["inputs_sub"])
        lines = []
        if docs:
            lines.append(
                L["docs"].format(
                    n=len(docs), names=", ".join(_cut(d.filename, 40) for d in docs[:4])
                )
            )
        if reqs:
            lines.append(L["reqs"].format(n=reqs))
        for a, p in data[:3]:
            pii = sum(c.pii_suspect for c in p.columns)
            line = L["data"].format(
                name=_cut(a.filename, 30),
                rows=f"{p.rows:,}",
                cols=len(p.columns),
                hint=p.readiness_hint,
            )
            lines.append(line + (L["pii"].format(n=pii) if pii else ""))
        for a in code[:2]:
            cp = a.code_profile
            lines.append(
                L["code"].format(
                    name=_cut(a.filename, 30),
                    lines=f"{cp.total_lines:,}",
                    langs=", ".join(list(cp.languages)[:4]),
                    fw=", ".join(cp.frameworks) or "-",
                )
            )
        self.text(
            s, Inches(0.75), Inches(1.5), Inches(11.8), Inches(3), lines, size=16, bullets=True
        )
        issues = [t(i) for _, p in data for i in p.issues][:5]
        if issues:
            self.text(
                s,
                Inches(0.75),
                Inches(4.6),
                Inches(11.8),
                Inches(0.4),
                L["data_notes"],
                size=15,
                bold=True,
                color=self.accent,
            )
            self.text(
                s,
                Inches(0.75),
                Inches(5.05),
                Inches(11.8),
                Inches(1.8),
                issues,
                size=13,
                bullets=True,
                color=GRAY,
            )

    def pattern_slide(self) -> None:
        p, L, t = self.run.pattern, self.L, self.t
        if not p:
            return
        s = self.slide(L["pattern"])
        no_ai = p.pattern.value == "no_ai_rule_based"
        self.rounded(
            s,
            Inches(0.75),
            Inches(1.35),
            Inches(11.8),
            Inches(1.1),
            "1F9D6B" if no_ai else self.soft,
        )
        label = L["no_ai"] if no_ai else L["pattern_labels"][p.pattern.value]
        self.text(
            s,
            Inches(1.0),
            Inches(1.5),
            Inches(9),
            Inches(0.8),
            label,
            size=28,
            bold=True,
            color="FFFFFF" if no_ai else self.accent,
        )
        conf = f"{L['confidence']}: {L['confidence_labels'][p.confidence.value]}"
        self.text(
            s,
            Inches(9.6),
            Inches(1.62),
            Inches(2.8),
            Inches(0.6),
            conf,
            size=14,
            bold=True,
            color="FFFFFF" if no_ai else INK,
            align=PP_ALIGN.RIGHT,
        )
        self.text(
            s,
            Inches(0.75),
            Inches(2.7),
            Inches(11.8),
            Inches(1.4),
            _cut(t(p.rationale), 480),
            size=15,
        )
        if p.rejected:
            rows = [
                [L["pattern_labels"][r.pattern.value], _cut(t(r.reason), 140)]
                for r in p.rejected[:5]
            ]
            self.table(
                s, Inches(0.75), Inches(4.2), Inches(11.8), L["rejected"], rows, [3.6, 8.2], size=12
            )

    def architecture_slide(self) -> None:
        arch, L = self.run.architecture, self.L
        if not arch:
            return
        s = self.slide(
            L["architecture"], f"{L['deployment']}: {L['deployment_labels'][arch.deployment.value]}"
        )
        graph = graph_for(self.run)
        columns = layers(graph)[:6]
        if not columns:
            return
        left, top, width, height = Inches(0.75), Inches(1.6), Inches(11.8), Inches(4.9)
        col_w = width / len(columns)
        box_w = min(Inches(2.3), int(col_w * 0.78))
        shapes = {}
        for ci, nodes in enumerate(columns):
            nodes = nodes[:5]
            row_h = height / len(nodes)
            box_h = min(Inches(0.95), int(row_h * 0.7))
            for ri, node in enumerate(nodes):
                x = int(left + col_w * ci + (col_w - box_w) / 2)
                y = int(top + row_h * ri + (row_h - box_h) / 2)
                box = self.rect(
                    s,
                    x,
                    y,
                    box_w,
                    box_h,
                    self.soft,
                    line=self.accent,
                    shape=MSO_SHAPE.ROUNDED_RECTANGLE,
                )
                frame = box.text_frame
                frame.word_wrap = True
                frame.vertical_anchor = MSO_ANCHOR.MIDDLE
                para = frame.paragraphs[0]
                para.alignment = PP_ALIGN.CENTER
                run = para.add_run()
                run.text = _cut(self.t(graph.nodes[node]), 40)
                run.font.size, run.font.name, run.font.bold = self.pt(12), self.font, True
                run.font.color.rgb = _rgb(INK)
                shapes[node] = box
        for a, b in graph.edges:
            if a in shapes and b in shapes:
                src, dst = shapes[a], shapes[b]
                conn = s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, 0, 0, Emu(1), Emu(1))
                same_col = abs(src.left - dst.left) < Inches(0.1)
                conn.begin_connect(src, 2 if same_col else 3)
                conn.end_connect(dst, 0 if same_col else 1)
                conn.line.color.rgb = _rgb(GRAY)
                conn.line.width = Pt(1.5)
                tail = OxmlElement("a:tailEnd")
                tail.set("type", "triangle")
                conn.line._get_or_add_ln().append(tail)

    def feasibility_slide(self) -> None:
        f, L, t = self.run.feasibility, self.L, self.t
        if not f:
            return
        s = self.slide(L["feasibility"], f"{L['recommend']}: {L['go'][f.go_recommendation.value]}")
        values = [f.data_readiness, f.technical_feasibility, f.business_value]
        for i, (label, value) in enumerate(zip(L["scores"], values, strict=True)):
            x = Inches(0.75 + i * 4.0)
            self.rounded(s, x, Inches(1.45), Inches(3.75), Inches(1.25), self.soft, 0.12)
            self.text(
                s,
                x + Inches(0.25),
                Inches(1.5),
                Inches(1.6),
                Inches(1.1),
                f"{value}/5",
                size=34,
                bold=True,
                color=self.accent,
            )
            self.text(
                s,
                x + Inches(1.8),
                Inches(1.82),
                Inches(1.9),
                Inches(0.6),
                label,
                size=14,
                bold=True,
            )
        risks = f.risks[:6]
        if risks:
            rows = [
                [
                    str(r.severity),
                    L["risk"][r.category.value],
                    _cut(t(r.description), 110),
                    _cut(t(r.mitigation), 110),
                ]
                for r in risks
            ]
            table = self.table(
                s,
                Inches(0.75),
                Inches(3.0),
                Inches(11.8),
                L["risk_headers"],
                rows,
                [0.9, 1.7, 4.6, 4.6],
                size=11,
            )
            for i, r in enumerate(risks, start=1):
                cell = table.cell(i, 0)
                cell.fill.solid()
                cell.fill.fore_color.rgb = _rgb(
                    "D64545" if r.severity >= 4 else "E3A008" if r.severity == 3 else "1F9D6B"
                )
                cell.text_frame.paragraphs[0].runs[0].font.color.rgb = _rgb("FFFFFF")

    def effort_slide(self) -> None:
        arch, basis, L, t = self.run.architecture, self.run.effort_basis, self.L, self.t
        if not arch:
            return
        s = self.slide(L["effort"], L["effort_sub"])
        if basis:
            pattern = L["pattern_labels"][basis.pattern.value].split(" –")[0].split("（")[0]
            if basis.multipliers:
                factors = "  ×  ".join(f"{m.factor} {t(m.label)}" for m in basis.multipliers)
                formula = L["formula"].format(p=pattern, f=factors, x=basis.factor)
            else:
                formula = L["formula_none"].format(p=pattern)
            self.rounded(s, Inches(0.75), Inches(1.4), Inches(11.8), Inches(0.75), self.soft2, 0.2)
            self.text(
                s,
                Inches(1.0),
                Inches(1.52),
                Inches(11.4),
                Inches(0.5),
                formula,
                size=15,
                bold=True,
                color=self.strong,
            )
        rows = [
            [
                PHASE_LABELS[e.phase.value],
                f"{e.min_person_days}–{e.max_person_days}",
                _cut(", ".join(t(x) for x in e.team), 60),
                _cut("; ".join(t(x) for x in e.deliverables), 120),
            ]
            for e in arch.estimates
        ]
        self.table(
            s,
            Inches(0.75),
            Inches(2.45),
            Inches(11.8),
            L["effort_headers"],
            rows,
            [1.6, 1.6, 3.6, 5.0],
            size=13,
        )
        # What the scope does NOT include; kept inside the design area (bottom 7.0 in = footer).
        if arch.out_of_scope:
            self.rounded(s, Inches(0.75), Inches(6.2), Inches(11.8), Inches(0.7), self.soft, 0.08)
            self.text(
                s,
                Inches(0.95),
                Inches(6.26),
                Inches(11.4),
                Inches(0.58),
                f"{L['out_of_scope']}: " + "; ".join(_cut(t(x), 50) for x in arch.out_of_scope[:4]),
                size=11,
                bold=True,
                color=self.strong,
            )

    def wbs_slide(self) -> None:
        wbs, L, t = self.run.wbs, self.L, self.t
        if not wbs:
            return
        total = sum(task.person_days for task in wbs.tasks)
        s = self.slide(L["wbs"], L["wbs_sub"].format(n=len(wbs.tasks), d=total))
        tasks = wbs.tasks[:13]
        rows = [
            [
                k.id,
                PHASE_LABELS[k.phase.value],
                _cut(t(k.name), 70),
                _cut(t(k.role), 30),
                str(k.person_days),
            ]
            for k in tasks
        ]
        self.table(
            s,
            Inches(0.75),
            Inches(1.45),
            Inches(11.8),
            L["wbs_headers"],
            rows,
            [0.8, 1.5, 5.6, 2.6, 1.3],
            size=11,
        )
        if len(wbs.tasks) > len(tasks):
            self.text(
                s,
                Inches(0.75),
                Inches(6.55),
                Inches(11),
                Inches(0.3),
                L["more"].format(n=len(wbs.tasks) - len(tasks)),
                size=11,
                color=GRAY,
            )

    def timeline_slide(self) -> None:
        wbs, sched, L = self.run.wbs, self.run.schedule, self.L
        if not (wbs and sched and sched.total_days):
            return
        weeks = max(1, -(-sched.total_days // 5))
        s = self.slide(L["timeline"], L["timeline_sub"].format(w=weeks, d=sched.total_days))
        tasks = wbs.tasks[:16]
        by_id = {task.id: task for task in sched.tasks}
        label_w, left, right = Inches(3.4), Inches(0.75), Inches(12.6)
        area = right - (left + label_w)
        top, row_h = Inches(1.75), min(Inches(0.31), int(Inches(4.9) / max(1, len(tasks))))
        step = max(1, -(-weeks // 20))
        for wk in range(0, weeks, step):
            x = int(left + label_w + area * wk / weeks)
            self.text(
                s,
                x,
                Inches(1.35),
                Inches(0.7),
                Inches(0.3),
                L["week"].format(n=wk + 1),
                size=9,
                color=GRAY,
            )
        for i, task in enumerate(tasks):
            y = int(top + row_h * i)
            self.text(
                s,
                left,
                y - Inches(0.04),
                label_w - Inches(0.1),
                row_h,
                _cut(f"{task.id} {self.t(task.name)}", 42),
                size=10,
            )
            st = by_id[task.id]
            x = int(left + label_w + area * st.start_day / (weeks * 5))
            w = max(Inches(0.08), int(area * (st.end_day - st.start_day) / (weeks * 5)))
            self.rounded(
                s,
                x,
                y + Inches(0.04),
                w,
                row_h - Inches(0.08),
                self.phase_colors[task.phase.value],
                0.3,
            )
        for i, (phase, color) in enumerate(self.phase_colors.items()):
            x = Inches(0.75 + i * 1.8)
            self.rect(s, x, Inches(6.67), Inches(0.3), Inches(0.16), color)
            self.text(
                s,
                x + Inches(0.38),
                Inches(6.6),
                Inches(1.4),
                Inches(0.3),
                PHASE_LABELS[phase],
                size=10,
                color=GRAY,
            )

    def requirements_slide(self) -> None:
        matrix, L, t = self.run.requirements, self.L, self.t
        if not matrix or matrix.skipped:
            return
        counts = coverage_counts(self.run)
        s = self.slide(L["requirements"], L["requirements_sub"].format(n=len(matrix.items)))
        for i, key in enumerate(["full", "partial", "not_supported", "needs_clarification"]):
            x = Inches(0.75 + i * 2.97)
            self.rounded(s, x, Inches(1.45), Inches(2.8), Inches(1.05), COVERAGE_COLORS[key], 0.12)
            self.text(
                s,
                x + Inches(0.2),
                Inches(1.5),
                Inches(1.0),
                Inches(0.9),
                str(counts[key]),
                size=30,
                bold=True,
                color="FFFFFF",
            )
            self.text(
                s,
                x + Inches(1.15),
                Inches(1.75),
                Inches(1.6),
                Inches(0.6),
                L["coverage"][key],
                size=13,
                bold=True,
                color="FFFFFF",
            )
        texts = requirement_texts(self.run)
        attention = [i for i in matrix.items if i.coverage != "full"][:8] or matrix.items[:8]
        rows = [
            [
                i.req_id,
                _cut(texts.get(i.req_id, ("",))[0], 80),
                L["coverage"][i.coverage],
                _cut(t(i.note), 70),
            ]
            for i in attention
        ]
        self.table(
            s,
            Inches(0.75),
            Inches(2.8),
            Inches(11.8),
            L["req_headers"],
            rows,
            [1.2, 5.0, 1.9, 3.7],
            size=11,
        )

    def pricing_slide(self) -> None:
        q, L, t = self.run.quotation, self.L, self.t
        if not q:
            return
        money = lambda v: format_money(v, q.currency)  # noqa: E731
        model = q.contract_model
        s = self.slide(L["pricing"], f"{L['contract'][model]} · {L['pricing_sub']} · {L['vat']}")
        if model == "odc":
            rows = [[t(m.role_label), f"{m.fte:g}", money(m.monthly_cost)] for m in q.odc_team]
            rows.append(
                [
                    L["odc_total"],
                    f"{sum(m.fte for m in q.odc_team):g}",
                    money(q.odc_monthly_cost or 0),
                ]
            )
            self.table(s, Inches(0.75), Inches(1.45), Inches(7.6), L["odc_headers"], rows, [4.0, 1.2, 2.4],
                       size=12, bold_last=True)  # fmt: skip
            self.text(s, Inches(0.75), Inches(1.6) + Inches(0.4) * (len(rows) + 1), Inches(7.6), Inches(0.5),
                      L["odc_period"].format(m=q.months, total=money(q.total)), size=15, bold=True,
                      color=self.accent)  # fmt: skip
        else:
            rows = [
                [
                    PHASE_LABELS[p.phase.value],
                    str(p.person_days),
                    money(p.amount),
                    f"{money(p.min_amount)} – {money(p.max_amount)}",
                ]
                for p in q.phases
            ]
            if model == "fixed_price":
                rows.append(
                    [
                        L["contingency"].format(p=f"{q.contingency_pct:g}"),
                        "",
                        money(q.contingency),
                        "",
                    ]
                )
            rows.append(
                [
                    L["total"] if model == "fixed_price" else L["tm_total"],
                    str(sum(p.person_days for p in q.phases)),
                    money(q.total),
                    f"{money(q.total_min)} – {money(q.total_max)}",
                ]
            )
            self.table(s, Inches(0.75), Inches(1.45), Inches(7.6), L["price_headers"], rows,
                       [2.4, 1.3, 1.9, 2.0], size=12, bold_last=True)  # fmt: skip
        if q.overhead_person_days or q.onsite_ratio:
            self.text(s, Inches(0.75), Inches(4.75), Inches(7.6), Inches(0.4),
                      L["overhead_note"].format(o=q.overhead_person_days, r=f"{q.onsite_ratio:g}"),
                      size=11, color=GRAY)  # fmt: skip
        y = Inches(1.45)
        if q.monthly_run_cost is not None:
            self.rounded(s, Inches(8.7), y, Inches(3.85), Inches(1.1), self.soft, 0.12)
            self.text(
                s,
                Inches(8.9),
                y + Inches(0.1),
                Inches(3.5),
                Inches(0.4),
                L["monthly"],
                size=12,
                color=GRAY,
            )
            self.text(s, Inches(8.9), y + Inches(0.45), Inches(3.5), Inches(0.5), money(q.monthly_run_cost),
                      size=20, bold=True, color=self.accent)  # fmt: skip
            y += Inches(1.3)
        self.text(s, Inches(8.7), y, Inches(3.85), Inches(0.4), L["milestones"], size=14, bold=True,
                  color=self.accent)  # fmt: skip
        shown = q.milestones[:6]
        lines = [f"{t(m.name)}: {m.percent:g}% · {money(m.amount)}" for m in shown]
        if len(q.milestones) > len(shown):
            lines.append("…")
        self.text(
            s,
            Inches(8.7),
            y + Inches(0.45),
            Inches(3.85),
            Inches(2.0),
            lines,
            size=12,
            bullets=True,
        )
        self.text(s, Inches(0.75), Inches(5.3), Inches(11.8), Inches(1.5), [t(a) for a in q.assumptions[:4]],
                  size=10, bullets=True, color=GRAY)  # fmt: skip

    def block_slides(self, position: str) -> None:
        """Standard company content (Cài đặt → Thư viện nội dung), inserted verbatim. Long blocks
        get their own slide; short ones are grouped as cards, up to four per slide."""
        blocks = render_blocks(self.run, self.kit.company, self.kit.blocks, position, self.lang)
        short = [b for b in blocks if len(b[1]) <= 420 and position == "end"]
        for title, body in [b for b in blocks if b not in short]:
            s = self.slide(title)
            lines = [
                line.strip().lstrip("-• ").strip() for line in body.splitlines() if line.strip()
            ]
            self.rect(
                s,
                Inches(0.75),
                Inches(1.5),
                Inches(0.08),
                Inches(0.5 * min(len(lines), 8)),
                self.accent,
            )
            self.text(
                s, Inches(1.05), Inches(1.4), Inches(11.4), Inches(5.2), lines, size=17, space=14
            )
        for page in range(0, len(short), 4):
            group = short[page : page + 4]
            s = self.slide(group[0][0] if len(group) == 1 else self.L["commitments"])
            cols = 1 if len(group) == 1 else 2
            rows = -(-len(group) // cols)
            w, h = 11.8 / cols - 0.25, 5.3 / rows - 0.25
            for i, (title, body) in enumerate(group):
                x, y = (
                    Inches(0.75 + (i % cols) * (w + 0.25)),
                    Inches(1.45 + (i // cols) * (h + 0.25)),
                )
                self.rounded(s, x, y, Inches(w), Inches(h), self.soft2, 0.06)
                self.rect(s, x, y, Inches(0.07), Inches(h), self.accent)
                self.text(s, x + Inches(0.3), y + Inches(0.2), Inches(w - 0.5), Inches(0.45), title, size=16,
                          bold=True, color=self.accent)  # fmt: skip
                lines = [
                    line.strip().lstrip("-• ").strip() for line in body.splitlines() if line.strip()
                ]
                self.text(s, x + Inches(0.3), y + Inches(0.75), Inches(w - 0.5), Inches(h - 0.9), lines,
                          size=12, space=6)  # fmt: skip

    def team_slide(self) -> None:
        """Project team derived from this run: WBS roles with their main tasks and person-days,
        plus the company overhead roles (PM, BrSE) priced in the quotation. Company profile and
        staff credentials belong in the content library, not here."""
        wbs, q, L, t = self.run.wbs, self.run.quotation, self.L, self.t
        if not wbs:
            return
        by_role: dict[str, list] = {}
        for task in wbs.tasks:
            by_role.setdefault(task.role, []).append(task)
        ranked = sorted(by_role.items(), key=lambda kv: -sum(k.person_days for k in kv[1]))
        rows = []
        for role, tasks in ranked[:6]:
            main = sorted(tasks, key=lambda k: -k.person_days)[:3]
            rows.append([
                _cut(t(role), 40),
                _cut("; ".join(t(k.name) for k in main), 110),
                str(sum(k.person_days for k in tasks)),
            ])  # fmt: skip
        overhead: dict[str, list] = {}
        for line in q.lines if q else []:
            if line.kind == "overhead":
                overhead.setdefault(line.role_key, [line.role_label, 0])[1] += line.person_days
        for key, (label, days) in overhead.items():
            rows.append([label, L["team_overhead"].get(key, ""), str(days)])
        self.table(
            self.slide(L["team"], L["team_sub"]),
            Inches(0.75),
            Inches(1.45),
            Inches(11.8),
            L["team_headers"],
            rows,
            [3.0, 7.0, 1.8],
            size=11,
        )

    def case_studies_slide(self) -> None:
        L, t = self.L, self.t
        matches = match_case_studies(self.run, self.kit.case_studies)
        if not matches:
            return
        s = self.slide(L["case_studies"], L["cs_sub"])
        width = 11.8 / len(matches)
        for i, match in enumerate(matches):
            cs = match["case"]
            x = Inches(0.75 + i * width)
            w = Inches(width - 0.25)
            self.rounded(s, x, Inches(1.45), w, Inches(5.2), self.soft2, 0.05)
            self.rect(s, x, Inches(1.45), w, Inches(0.08), self.accent)
            meta = " · ".join(
                str(v) for v in (t(cs.industry), L["markets"][cs.market], cs.year) if v
            )
            self.text(s, x + Inches(0.2), Inches(1.65), w - Inches(0.4), Inches(0.9), _cut(t(cs.title), 90),
                      size=14, bold=True)  # fmt: skip
            self.text(
                s,
                x + Inches(0.2),
                Inches(2.5),
                w - Inches(0.4),
                Inches(0.3),
                meta,
                size=10,
                color=GRAY,
            )
            blocks = [
                (L["cs_challenge"], [_cut(t(cs.challenge), 150)]),
                (L["cs_solution"], [_cut(t(cs.solution), 150)]),
                (L["cs_results"], [_cut(t(r), 70) for r in cs.results[:3]]),
            ]
            y = 2.9
            for label, lines in blocks:
                self.text(s, x + Inches(0.2), Inches(y), w - Inches(0.4), Inches(0.3), label, size=10, bold=True,
                          color=self.accent)  # fmt: skip
                self.text(s, x + Inches(0.2), Inches(y + 0.28), w - Inches(0.4), Inches(1.0), lines, size=10,
                          bullets=label == L["cs_results"], space=2)  # fmt: skip
                y += 1.2

    def assumptions_slide(self) -> None:
        L, t = self.L, self.t
        assumptions = [t(a) for a in self.run.pattern.assumptions] if self.run.pattern else []
        questions = [q.question for q in self.run.gaps.questions] if self.run.gaps else []
        if not (assumptions or questions):
            return
        s = self.slide(L["assumptions"])
        self.text(
            s,
            Inches(0.75),
            Inches(1.4),
            Inches(5.6),
            Inches(0.4),
            L["assumptions_col"],
            size=17,
            bold=True,
            color=self.accent,
        )
        self.text(
            s,
            Inches(0.75),
            Inches(1.9),
            Inches(5.6),
            Inches(4.8),
            [_cut(a, 140) for a in assumptions[:8]] or [L["none"]],
            size=13,
            bullets=True,
        )
        self.text(
            s,
            Inches(6.9),
            Inches(1.4),
            Inches(5.6),
            Inches(0.4),
            L["questions_col"],
            size=17,
            bold=True,
            color=self.accent,
        )
        self.text(
            s,
            Inches(6.9),
            Inches(1.9),
            Inches(5.6),
            Inches(4.8),
            [_cut(q, 140) for q in questions[:8]] or [L["none"]],
            size=13,
            bullets=True,
        )

    def next_steps_slide(self) -> None:
        L = self.L
        go = self.run.feasibility.go_recommendation.value if self.run.feasibility else "go_with_poc"
        s = self.slide(L["next"], L["go"][go])
        for i, step in enumerate(L["next_steps"][go]):
            y = Inches(1.7 + i * 1.5)
            circle = self.rect(
                s, Inches(0.9), y, Inches(0.8), Inches(0.8), self.accent, shape=MSO_SHAPE.OVAL
            )
            circle.text_frame.paragraphs[0].alignment = PP_ALIGN.CENTER
            run = circle.text_frame.paragraphs[0].add_run()
            run.text, run.font.size, run.font.bold, run.font.name = (
                str(i + 1),
                self.pt(22),
                True,
                self.font,
            )
            run.font.color.rgb = _rgb("FFFFFF")
            self.text(
                s, Inches(2.0), y + Inches(0.15), Inches(10), Inches(0.6), step, size=20, bold=True
            )

    def build(self) -> bytes:
        for make in (
            self.title_slide, lambda: self.block_slides("start"), self.context_slide, self.inputs_slide, self.pattern_slide, self.architecture_slide,
            self.feasibility_slide, self.effort_slide, self.wbs_slide, self.timeline_slide,
            self.team_slide, self.requirements_slide, self.case_studies_slide,
            self.pricing_slide, lambda: self.block_slides("end"),
            self.assumptions_slide, self.next_steps_slide,
        ):  # fmt: skip
            make()
        buffer = io.BytesIO()
        self.prs.save(buffer)
        return buffer.getvalue()


def build_slides(
    run: ScopingRun,
    lang: str = "vi",
    tr: Callable[[str], str] | None = None,
    kit: ExportKit | None = None,
) -> bytes:
    return Deck(run, lang, tr, kit).build()
