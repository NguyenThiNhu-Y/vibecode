"""Fill company templates (.pptx / .docx / .xlsx): {{placeholder}} text and content markers.

Conventions a template designer follows (shown in Cài đặt → Template):
- Placeholders anywhere in text: {{company_name}}, {{company_short}}, {{client_name}},
  {{project_name}}, {{proposal_title}}, {{date}}, {{version}}, {{confidential_footer}}.
- PowerPoint: the content layout may contain a shape whose text is {{SCOPEAI_CONTENT}}; its box is
  where ScopeAI draws tables and diagrams (the marker itself is removed).
- Word: a paragraph {{SCOPEAI_BODY}} marks where the proposal body goes.
- {{company_logo}}: a PowerPoint shape (master/layout) whose whole text is the marker gets the logo
  uploaded in Cài đặt → Công ty fitted inside it (the shape stays as the logo's card); in Word the
  marker becomes the logo inline. Without an uploaded logo it falls back to {{company_short}}.
"""

import copy
import re
from dataclasses import dataclass
from pathlib import Path
from typing import Any

from docx.shared import Cm
from docx.text.run import Run
from lxml import etree
from pptx import Presentation
from pptx.enum.shapes import PP_PLACEHOLDER
from pptx.opc.constants import RELATIONSHIP_TYPE as RT
from pptx.util import Emu, Inches

from app.company import fill

CONTENT_MARKER = "{{SCOPEAI_CONTENT}}"
BODY_MARKER = "{{SCOPEAI_BODY}}"
LOGO_MARKER = "{{company_logo}}"
A_NS = "http://schemas.openxmlformats.org/drawingml/2006/main"
TITLE_TYPES = {PP_PLACEHOLDER.TITLE, PP_PLACEHOLDER.CENTER_TITLE}
KEEP_ON_SLIDE = {PP_PLACEHOLDER.SLIDE_NUMBER, PP_PLACEHOLDER.FOOTER}


# ---------- colors ----------
def mix(hex_color: str, other: str, amount: float) -> str:
    """Blend hex_color towards other by amount (0..1)."""
    a = [int(hex_color[i : i + 2], 16) for i in (0, 2, 4)]
    b = [int(other[i : i + 2], 16) for i in (0, 2, 4)]
    return "".join(f"{round(x + (y - x) * amount):02X}" for x, y in zip(a, b, strict=True))


def theme_accent(prs: Any) -> str | None:
    """accent1 of the template theme, so tables and diagrams follow the company colors."""
    try:
        theme = prs.slide_master.part.part_related_by(RT.THEME)
        root = etree.fromstring(theme.blob)
    except Exception:
        return None
    accent = root.find(f".//{{{A_NS}}}accent1")
    if accent is None or len(accent) == 0:
        return None
    color = accent[0]
    value = color.get("val") if color.tag.endswith("srgbClr") else color.get("lastClr")
    return value.upper() if value and re.fullmatch(r"[0-9A-Fa-f]{6}", value) else None


# ---------- text replacement ----------
def _fill_paragraph(paragraph: Any, values: dict[str, str]) -> None:
    runs = paragraph.runs
    full = "".join(r.text for r in runs)
    if "{{" not in full:
        return
    new = fill(full, values)
    if new != full:
        runs[0].text = new
        for r in runs[1:]:
            r.text = ""


def fill_shapes(shapes: Any, values: dict[str, str]) -> None:
    for shape in shapes:
        if shape.shape_type is not None and shape.shape_type == 6:  # group
            fill_shapes(shape.shapes, values)
            continue
        if getattr(shape, "has_text_frame", False) and shape.has_text_frame:
            for p in shape.text_frame.paragraphs:
                _fill_paragraph(p, values)
        if getattr(shape, "has_table", False) and shape.has_table:
            for row in shape.table.rows:
                for cell in row.cells:
                    for p in cell.text_frame.paragraphs:
                        _fill_paragraph(p, values)


def fit_logo(
    left: int,
    top: int,
    width: int,
    height: int,
    logo: Path,
    pad: float = 0.12,
    align: str = "center",
) -> tuple[int, int, int, int]:
    """Largest box with the logo's aspect ratio inside (left, top, width, height), EMU."""
    from PIL import Image

    with Image.open(logo) as img:
        img_w, img_h = img.size
    gap = int(min(width, height) * pad)
    scale = min((width - 2 * gap) / img_w, (height - 2 * gap) / img_h)
    cx, cy = int(img_w * scale), int(img_h * scale)
    x = {"left": left + gap, "right": left + width - gap - cx}.get(align, left + (width - cx) // 2)
    return x, top + (height - cy) // 2, cx, cy


def place_pptx_logo(shapes: Any, part: Any, logo: Path) -> None:
    """Put the logo inside every {{company_logo}} shape of a master/layout (marker text cleared)."""
    for shape in list(shapes):
        if not getattr(shape, "has_text_frame", False) or not shape.has_text_frame:
            continue
        if shape.text_frame.text.strip() != LOGO_MARKER:
            continue
        for paragraph in shape.text_frame.paragraphs:
            for run in paragraph.runs:
                run.text = ""
        _, rid = part.get_or_add_image_part(str(logo))
        x, y, cx, cy = fit_logo(shape.left, shape.top, shape.width, shape.height, logo)
        pic = shapes._spTree.add_pic(shapes._next_shape_id, "Logo", "", rid, x, y, cx, cy)
        shape._element.addnext(pic)  # right above its card


def drop_slides(prs: Any) -> None:
    """Remove the example slides a template ships with (their parts are not saved)."""
    ids = prs.slides._sldIdLst
    for sld_id in list(ids):
        prs.part.drop_rel(sld_id.rId)
        ids.remove(sld_id)


# ---------- PowerPoint ----------
@dataclass
class PptxTemplate:
    prs: Any
    cover_layout: Any
    content_layout: Any
    region: tuple[int, int, int, int]  # left, top, width, height (EMU) for the slide body
    accent: str | None


def _layout_by_name(prs: Any, name: str | None) -> Any | None:
    for layout in prs.slide_layouts:
        if name and layout.name.strip().lower() == name.strip().lower():
            return layout
    return None


def _title_placeholder(layout: Any) -> Any | None:
    return next((p for p in layout.placeholders if p.placeholder_format.type in TITLE_TYPES), None)


def find_layouts(prs: Any, content_name: str | None) -> tuple[Any, Any]:
    layouts = list(prs.slide_layouts)
    cover = _layout_by_name(prs, "Title Slide") or next(
        (
            lo
            for lo in layouts
            if any(
                p.placeholder_format.type == PP_PLACEHOLDER.CENTER_TITLE for p in lo.placeholders
            )
        ),
        layouts[0],
    )
    content = _layout_by_name(prs, content_name) or _layout_by_name(prs, "Title Only")
    if content is None:  # the layout with a title and the fewest other placeholders
        candidates = [
            lo for lo in layouts if _title_placeholder(lo) is not None and lo is not cover
        ]
        content = (
            min(candidates, key=lambda lo: len(lo.placeholders)) if candidates else layouts[-1]
        )
    return cover, content


def _marker_box(shapes: Any) -> tuple[int, int, int, int] | None:
    for shape in shapes:
        if shape.has_text_frame and shape.text_frame.text.strip() == CONTENT_MARKER:
            box = (shape.left, shape.top, shape.width, shape.height)
            shape.element.getparent().remove(shape.element)
            return box
    return None


def open_pptx_template(
    path: Path, cfg: dict[str, Any], values: dict[str, str], logo: Path | None = None
) -> PptxTemplate:
    prs = Presentation(str(path))
    drop_slides(prs)
    if logo:
        place_pptx_logo(prs.slide_master.shapes, prs.slide_master.part, logo)
        for layout in prs.slide_layouts:
            place_pptx_logo(layout.shapes, layout.part, logo)
    fill_shapes(prs.slide_master.shapes, values)
    for layout in prs.slide_layouts:
        fill_shapes(layout.shapes, values)
    cover, content = find_layouts(prs, str(cfg.get("content_layout") or "") or None)
    width, height = prs.slide_width, prs.slide_height
    region = None
    for layout in prs.slide_layouts:  # the marker must never show, whichever layout holds it
        box = _marker_box(layout.shapes)
        if layout is content and box:
            region = box
    master_box = _marker_box(prs.slide_master.shapes)
    region = region or master_box
    if region is None:
        title = _title_placeholder(content)
        margin = Inches(0.5)
        top = (title.top + title.height + Inches(0.12)) if title is not None else Inches(1.2)
        left = title.left if title is not None and title.left < width / 3 else margin
        bottom = height - Emu(int(Inches(float(cfg.get("bottom_margin_in", 0.55)))))
        region = (left, top, width - left - margin, max(Inches(2), bottom - top))
    return PptxTemplate(prs, cover, content, region, theme_accent(prs))


def tidy_placeholders(slide: Any, layout: Any) -> None:
    """Call after filling title/subtitle. Drops empty placeholders PowerPoint would show as
    'Click to add text'; copies the layout's footer / slide number so numbering follows it."""
    for ph in list(slide.placeholders):
        if not ph.has_text_frame or not ph.text_frame.text.strip():
            ph.element.getparent().remove(ph.element)
    tree = slide.shapes._spTree
    for ph in layout.placeholders:
        if ph.placeholder_format.type in KEEP_ON_SLIDE:
            tree.append(copy.deepcopy(ph.element))


# ---------- Word ----------
def _paragraphs(container: Any) -> Any:
    yield from container.paragraphs
    for table in container.tables:
        for row in table.rows:
            for cell in row.cells:
                yield from _paragraphs(cell)


def _docx_paragraphs(doc: Any) -> Any:
    """Body paragraphs (tables included), then those of every own header/footer."""
    yield from _paragraphs(doc)
    for section in doc.sections:
        for part in (
            section.header,
            section.footer,
            section.first_page_header,
            section.first_page_footer,
        ):
            if part is not None and not part.is_linked_to_previous:
                yield from _paragraphs(part)


def _docx_logo(paragraph: Any, logo: Path) -> None:
    """Split the run holding {{company_logo}} into text before / logo / text after; the three
    runs keep the original run's formatting."""
    runs = paragraph.runs
    run = next((r for r in runs if LOGO_MARKER in r.text), None)
    if run is None:  # marker split across runs: merge them first
        full = paragraph.text
        runs[0].text = full
        for r in runs[1:]:
            r.text = ""
        run = runs[0]
    before, _, after = run.text.partition(LOGO_MARKER)
    picture, tail = copy.deepcopy(run._r), copy.deepcopy(run._r)
    run.text = before
    run._r.addnext(picture)
    picture.addnext(tail)
    Run(tail, paragraph).text = after
    image = Run(picture, paragraph)
    image.text = ""
    image.add_picture(str(logo), height=Cm(1.0))


def fill_docx(doc: Any, values: dict[str, str], logo: Path | None = None) -> None:
    # With a logo the marker is left for _docx_logo; filling text first, because merging runs
    # (_fill_paragraph) would wipe a picture already placed in them.
    text_values = {k: v for k, v in values.items() if not (logo and k == "company_logo")}
    for p in list(_docx_paragraphs(doc)):
        _fill_paragraph(p, text_values)
        if logo and LOGO_MARKER in p.text:
            _docx_logo(p, logo)


def find_body_marker(doc: Any) -> Any | None:
    return next((p for p in doc.paragraphs if p.text.strip() == BODY_MARKER), None)


def move_after_marker(doc: Any, marker: Any, first_new: int) -> None:
    """Move body elements appended from index first_new to where the marker paragraph was."""
    body = doc.element.body
    new = [el for el in list(body)[first_new:] if not el.tag.endswith("}sectPr")]
    anchor = marker._p
    for el in new:
        anchor.addnext(el)
        anchor = el
    marker._p.getparent().remove(marker._p)


# ---------- Excel ----------
def fill_workbook(wb: Any, values: dict[str, str]) -> None:
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for cell in row:
                if (
                    isinstance(cell.value, str)
                    and "{{" in cell.value
                    and not cell.value.startswith("=")
                ):
                    cell.value = fill(cell.value, values)
