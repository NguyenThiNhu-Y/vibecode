"""Excel workbook: overview, effort (live formulas), WBS, timeline, requirements, risks, data."""

import io

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from app.company import placeholders
from app.exports.common import (
    ACCENT,
    COVERAGE_COLORS,
    COVERAGE_LABELS,
    DEPLOYMENT_LABELS,
    GO_LABELS,
    PATTERN_LABELS,
    PHASE_LABELS,
    RISK_LABELS,
    requirement_texts,
)
from app.exports.plan import schedule_for, task_rows, week_of, week_starts
from app.exports.templating import fill_workbook
from app.schemas.run import ScopingRun
from app.schemas.schedule import ScheduleResult
from app.schemas.settings import SheetMapping
from app.schemas.wbs import ordered
from app.templates_store import ExportKit

HEADER_FILL = PatternFill("solid", fgColor=ACCENT)
HEADER_FONT = Font(bold=True, color="FFFFFF")
THIN = Side(style="thin", color="D9D9D9")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
WRAP = Alignment(wrap_text=True, vertical="top")


def _header(ws: Worksheet, row: int, titles: list[str], widths: list[int] | None = None) -> None:
    for col, title in enumerate(titles, start=1):
        cell = ws.cell(row=row, column=col, value=title)
        cell.fill, cell.font, cell.border = HEADER_FILL, HEADER_FONT, BORDER
        cell.alignment = Alignment(vertical="center", wrap_text=True)
        if widths:
            ws.column_dimensions[get_column_letter(col)].width = widths[col - 1]
    ws.freeze_panes = ws.cell(row=row + 1, column=1)


def _row(ws: Worksheet, row: int, values: list) -> None:
    for col, value in enumerate(values, start=1):
        cell = ws.cell(row=row, column=col, value=value)
        cell.border, cell.alignment = BORDER, WRAP


def _title(ws: Worksheet, text: str) -> None:
    ws["A1"] = text
    ws["A1"].font = Font(bold=True, size=14, color=ACCENT)


def _duration(schedule: ScheduleResult | None) -> str:
    if not schedule:
        return ""
    start, end = schedule.phases[0].start, schedule.phases[-1].end
    days = sum(p.working_days for p in schedule.phases)
    return f"{start:%d/%m/%Y} → {end:%d/%m/%Y}: {days} ngày làm việc (~{-(-days // 5)} tuần)"


def _overview(ws: Worksheet, run: ScopingRun) -> None:
    ws.title = "Tổng quan"
    _title(ws, run.proposal.title if run.proposal else "ScopeAI – Proposal sơ bộ")
    rows = [
        ("Mã phiên", run.id),
        ("Mục tiêu", run.intake.business_goal if run.intake else ""),
        (
            "Hướng giải pháp",
            PATTERN_LABELS.get(run.pattern.pattern.value, "") if run.pattern else "",
        ),
        (
            "Khuyến nghị",
            GO_LABELS[run.feasibility.go_recommendation.value] if run.feasibility else "",
        ),
        (
            "Triển khai",
            DEPLOYMENT_LABELS[run.architecture.deployment.value] if run.architecture else "",
        ),
        ("Tổng thời gian dự kiến", _duration(schedule_for(run))),
        ("Trạng thái review", run.status.value),
        ("Ghi chú", "Bản nháp do ScopeAI tạo; số liệu effort tính bằng code từ bảng chuẩn."),
    ]
    for i, (k, v) in enumerate(rows, start=3):
        ws.cell(row=i, column=1, value=k).font = Font(bold=True)
        ws.cell(row=i, column=2, value=v).alignment = WRAP
    ws.column_dimensions["A"].width = 26
    ws.column_dimensions["B"].width = 90


def _effort(ws: Worksheet, run: ScopingRun) -> None:
    _title(ws, "Effort (ngày công) – công thức tính bằng code")
    basis = run.effort_basis
    ws["A3"], ws["A3"].font = "Hệ số điều chỉnh", Font(bold=True)
    multipliers = basis.multipliers if basis else []
    for i, m in enumerate(multipliers):
        ws.cell(row=4 + i, column=1, value=m.label)
        ws.cell(row=4 + i, column=2, value=m.factor)
    factor_row = 4 + len(multipliers)
    ws.cell(row=factor_row, column=1, value="Tích hệ số").font = Font(bold=True)
    ws.cell(
        row=factor_row,
        column=2,
        value=f"=PRODUCT(B4:B{factor_row - 1})" if multipliers else 1,
    ).font = Font(bold=True, color=ACCENT)
    factor_ref = f"$B${factor_row}"

    start = factor_row + 2
    _header(
        ws,
        start,
        [
            "Giai đoạn",
            "Chuẩn min",
            "Chuẩn max",
            "Gốc min",
            "Gốc max",
            "Đề xuất min",
            "Đề xuất max",
            "Ghi chú điều chỉnh",
        ],
        [30, 12, 12, 12, 12, 13, 13, 50],
    )
    for i, estimate in enumerate(run.architecture.estimates if run.architecture else [], start=1):
        r = start + i
        base = (basis.base.get(estimate.phase) if basis else None) or [None, None]
        _row(
            ws,
            r,
            [
                PHASE_LABELS[estimate.phase.value],
                base[0],
                base[1],
                f"=ROUND(B{r}*{factor_ref},0)" if base[0] is not None else None,
                f"=ROUND(C{r}*{factor_ref},0)" if base[1] is not None else None,
                estimate.min_person_days,
                estimate.max_person_days,
                estimate.adjustment_note or "",
            ],
        )


def _wbs(ws: Worksheet, run: ScopingRun) -> None:
    """Whole WBS tree; the detailed, formula-driven version is the Bidding workbook."""
    _header(
        ws,
        1,
        ["ID", "Giai đoạn", "Cấp", "Đầu việc", "Loại", "Ưu tiên", "Ngày công", "Phụ thuộc",
         "Bàn giao"],
        [8, 12, 6, 50, 9, 9, 11, 14, 36],
    )  # fmt: skip
    if not run.wbs:
        return
    items = ordered(run.wbs)
    for i, item in enumerate(items, start=2):
        _row(
            ws,
            i,
            [
                item.id,
                PHASE_LABELS[item.phase.value],
                item.level,
                ("    " * (item.level - 1)) + item.name,
                item.type.value if item.type else "",
                item.priority.value,
                item.estimate_md,
                ", ".join(item.depends_on),
                item.deliverable or "",
            ],
        )
        if item.level == 1:
            ws.cell(row=i, column=4).font = Font(bold=True)
    last = len(items) + 1
    r = last + 2
    for label in PHASE_LABELS.values():  # level-1 rows only: parents already hold their sums
        ws.cell(row=r, column=4, value=f"Tổng {label}").font = Font(bold=True)
        ws.cell(
            row=r, column=7, value=f'=SUMIFS(G2:G{last},B2:B{last},"{label}",C2:C{last},1)'
        ).font = Font(bold=True)
        r += 1
    ws.cell(row=r, column=4, value="Tổng cộng").font = Font(bold=True, color=ACCENT)
    ws.cell(row=r, column=7, value=f"=SUMIFS(G2:G{last},C2:C{last},1)").font = Font(
        bold=True, color=ACCENT
    )


def _timeline(ws: Worksheet, run: ScopingRun) -> None:
    schedule = schedule_for(run)
    if not (run.wbs and schedule):
        return
    weeks = week_starts(schedule)
    _header(ws, 1, ["ID", "Đầu việc", "Bắt đầu", "Kết thúc"] + [w.strftime("%d/%m") for w in weeks])
    for col, width in zip("ABCD", (7, 42, 11, 11), strict=True):
        ws.column_dimensions[col].width = width
    fills = {
        "poc": PatternFill("solid", fgColor="F9A66C"),
        "mvp": PatternFill("solid", fgColor=ACCENT),
        "production": PatternFill("solid", fgColor="B5420E"),
    }
    rows = [t for t in task_rows(run) if t.week_from]
    for i, task in enumerate(rows, start=2):
        _row(ws, i, [task.id, task.name, task.start, task.end])
        for col in (3, 4):
            ws.cell(row=i, column=col).number_format = "dd/mm/yyyy"
        for w in range(task.week_from - 1, task.week_to):  # type: ignore[operator]
            ws.cell(row=i, column=5 + w).fill = fills[task.phase.value]
    r = len(rows) + 3
    for m in schedule.milestones:
        ws.cell(row=r, column=1, value=m.id).font = Font(bold=True)
        ws.cell(row=r, column=2, value=m.name).font = Font(bold=True)
        ws.cell(row=r, column=3, value=m.date).number_format = "dd/mm/yyyy"
        ws.cell(row=r, column=4 + week_of(schedule, m.date), value="◆")
        r += 1
    for w in range(len(weeks)):
        ws.column_dimensions[get_column_letter(5 + w)].width = 6.5


def _requirements(ws: Worksheet, run: ScopingRun) -> None:
    texts = requirement_texts(run)
    _header(
        ws,
        1,
        [
            "Mã",
            "Nhóm",
            "Ưu tiên",
            "Yêu cầu của khách",
            "Mức đáp ứng",
            "Thành phần đảm nhận",
            "Ghi chú",
        ],
        [12, 16, 10, 60, 18, 24, 50],
    )
    items = run.requirements.items if run.requirements else []
    for i, item in enumerate(items, start=2):
        text, priority, category = texts.get(item.req_id, ("", None, None))
        _row(
            ws,
            i,
            [
                item.req_id,
                category,
                priority,
                text,
                COVERAGE_LABELS[item.coverage],
                item.component,
                item.note,
            ],
        )
        cell = ws.cell(row=i, column=5)
        cell.fill = PatternFill("solid", fgColor=COVERAGE_COLORS[item.coverage])
        cell.font = Font(bold=True, color="FFFFFF")
    if items:
        ws.auto_filter.ref = f"A1:G{len(items) + 1}"
        r = len(items) + 3
        for label in COVERAGE_LABELS.values():
            ws.cell(row=r, column=4, value=label).font = Font(bold=True)
            ws.cell(row=r, column=5, value=f'=COUNTIF(E2:E{len(items) + 1},"{label}")')
            r += 1


def _quotation(ws: Worksheet, run: ScopingRun) -> None:
    """Price sheet with live formulas so the customer can see (and adjust) the calculation."""
    q = run.quotation
    if q is None:
        return
    model = {"fixed_price": "Trọn gói", "time_material": "T&M", "odc": "ODC"}[q.contract_model]
    _title(ws, f"Báo giá sơ bộ ({q.currency}, {model}) – số liệu giả lập, chưa gồm VAT")
    _header(
        ws,
        3,
        ["Giai đoạn", "Vai trò", "Ngày công", "Đơn giá / ngày", "Thành tiền", "Loại"],
        [16, 26, 12, 18, 20, 14],
    )
    money = "#,##0" if q.currency != "USD" else "#,##0.00"
    first = 4
    for i, line in enumerate(q.lines):
        r = first + i
        _row(
            ws,
            r,
            [
                PHASE_LABELS[line.phase.value],
                line.role_label,
                line.person_days,
                line.day_rate,
                f"=C{r}*D{r}",
                "Quản lý (overhead)" if line.kind == "overhead" else "WBS",
            ],
        )
        ws.cell(row=r, column=4).number_format = money
        ws.cell(row=r, column=5).number_format = money
    last = first + len(q.lines) - 1
    r = last + 2
    total_label = {
        "fixed_price": "TỔNG (đã gồm dự phòng)",
        "time_material": "TỔNG DỰ TOÁN (T&M)",
        "odc": "Dự toán theo WBS (tham khảo)",
    }[q.contract_model]
    rows = [
        ("Cộng chi phí nhân sự", f"=SUM(E{first}:E{last})"),
        ("% dự phòng rủi ro", q.contingency_pct / 100),
        ("Dự phòng rủi ro", f"=E{r}*E{r + 1}"),
        (total_label, f"=E{r}+E{r + 2}"),
    ]
    for offset, (label, value) in enumerate(rows):
        ws.cell(row=r + offset, column=4, value=label).font = Font(bold=True)
        cell = ws.cell(row=r + offset, column=5, value=value)
        cell.number_format = "0%" if offset == 1 else money
        cell.font = Font(bold=True, color=ACCENT if offset == 3 else "000000")
    total_ref = f"E{r + 3}"
    r += 5
    if q.contract_model == "odc":
        _header(ws, r, ["Đội ODC", "FTE", "", "", "Chi phí / tháng"])
        for i, m in enumerate(q.odc_team, start=1):
            _row(ws, r + i, [m.role_label, m.fte, "", "", m.monthly_cost])
            ws.cell(row=r + i, column=5).number_format = money
        end = r + len(q.odc_team)
        ws.cell(row=end + 1, column=4, value="Chi phí đội / tháng").font = Font(bold=True)
        ws.cell(row=end + 1, column=5, value=f"=SUM(E{r + 1}:E{end})").number_format = money
        ws.cell(row=end + 2, column=4, value="Số tháng").font = Font(bold=True)
        ws.cell(row=end + 2, column=5, value=q.months)
        ws.cell(row=end + 3, column=4, value="TỔNG ODC").font = Font(bold=True, color=ACCENT)
        cell = ws.cell(row=end + 3, column=5, value=f"=E{end + 1}*E{end + 2}")
        cell.number_format, cell.font = money, Font(bold=True, color=ACCENT)
        total_ref = f"E{end + 3}"
        r = end + 5
    ws.cell(row=r, column=4, value="Khoảng ước tính").font = Font(bold=True)
    ws.cell(row=r, column=5, value=f"{q.total_min:,.0f} – {q.total_max:,.0f}")
    if q.monthly_run_cost is not None:
        r += 1
        ws.cell(row=r, column=4, value="Vận hành / tháng (ước tính)").font = Font(bold=True)
        ws.cell(row=r, column=5, value=q.monthly_run_cost).number_format = money
    r += 2
    _header(ws, r, ["Mốc thanh toán", "", "%", "", "Số tiền"])
    ws.freeze_panes = "A4"
    for i, m in enumerate(q.milestones, start=1):
        _row(ws, r + i, [m.name, "", m.percent / 100, "", f"=ROUND({total_ref}*C{r + i},0)"])
        ws.cell(row=r + i, column=3).number_format = "0%"
        ws.cell(row=r + i, column=5).number_format = money
    r += len(q.milestones) + 2
    for i, a in enumerate(q.assumptions):
        ws.cell(row=r + i, column=1, value=f"• {a}")


def _risks(ws: Worksheet, run: ScopingRun) -> None:
    _header(ws, 1, ["Mức độ (1–5)", "Loại", "Rủi ro", "Cách giảm thiểu"], [13, 16, 60, 60])
    for i, risk in enumerate(run.feasibility.risks if run.feasibility else [], start=2):
        _row(
            ws,
            i,
            [risk.severity, RISK_LABELS[risk.category.value], risk.description, risk.mitigation],
        )


def _data(ws: Worksheet, run: ScopingRun) -> None:
    _header(
        ws,
        1,
        ["File", "Sheet", "Cột", "Kiểu", "% trống", "Giá trị khác nhau", "Nghi PII", "Ví dụ"],
        [24, 14, 26, 10, 10, 16, 10, 40],
    )
    r = 2
    for att in run.attachments:
        for profile in att.data_profiles:
            for col in profile.columns:
                _row(
                    ws,
                    r,
                    [
                        att.filename,
                        profile.sheet,
                        col.name,
                        col.dtype,
                        round(col.null_ratio * 100),
                        col.unique,
                        "Có" if col.pii_suspect else "",
                        ", ".join(col.examples),
                    ],
                )
                r += 1


def _target(wb: Workbook, mapping: SheetMapping):
    return wb[mapping.sheet] if mapping.sheet in wb.sheetnames else wb.create_sheet(mapping.sheet)


def _put(ws: Worksheet, mapping: SheetMapping, row: int, values: dict[str, object]) -> None:
    for key, value in values.items():
        col = mapping.columns.get(key)
        if col:
            cell = ws[f"{col}{row}"]
            cell.value = value
            cell.border, cell.alignment = BORDER, WRAP


def _fill_template(wb: Workbook, run: ScopingRun, kit: ExportKit) -> None:
    """Company estimate template: WBS and price lines written where the column mapping says."""
    cfg = kit.config.workbook
    if run.wbs:  # one row per level-2 task; "role" holds its work types
        mapping, ws = cfg["wbs"], _target(wb, cfg["wbs"])
        for i, task in enumerate(task_rows(run)):
            _put(ws, mapping, mapping.start_row + i, {
                "id": task.id, "phase": PHASE_LABELS[task.phase.value], "task": task.name,
                "role": ", ".join(task.types), "person_days": task.md,
                "start_week": task.week_from, "end_week": task.week_to,
            })  # fmt: skip
    if run.quotation:
        mapping, ws = cfg["pricing"], _target(wb, cfg["pricing"])
        cols = mapping.columns
        for i, line in enumerate(run.quotation.lines):
            r = mapping.start_row + i
            amount = (
                f"={cols['person_days']}{r}*{cols['day_rate']}{r}"
                if "person_days" in cols and "day_rate" in cols
                else line.amount
            )
            _put(ws, mapping, r, {
                "phase": PHASE_LABELS[line.phase.value], "role": line.role_label,
                "person_days": line.person_days, "day_rate": line.day_rate, "amount": amount,
            })  # fmt: skip
            for key in ("day_rate", "amount"):
                if key in cols:
                    ws[f"{cols[key]}{r}"].number_format = "#,##0"


def _sheet(wb: Workbook, name: str) -> Worksheet:
    """New ScopeAI sheet; never overwrites a sheet the company template already has."""
    return wb.create_sheet(name if name not in wb.sheetnames else f"{name} (ScopeAI)")


def build_workbook(run: ScopingRun, kit: ExportKit | None = None) -> bytes:
    kit = kit or ExportKit.load()
    template = kit.templates.get("workbook")
    if template:
        wb = load_workbook(str(template))
        fill_workbook(wb, placeholders(run, kit.company))
        _fill_template(wb, run, kit)
        _effort(_sheet(wb, "Effort"), run)
    else:
        wb = Workbook()
        _overview(wb.active, run)
        _effort(_sheet(wb, "Effort"), run)
    if run.quotation:
        _quotation(_sheet(wb, "Báo giá"), run)
    if not template:
        _wbs(_sheet(wb, "WBS"), run)
    if run.wbs:
        _timeline(_sheet(wb, "Timeline"), run)
    if run.requirements and not run.requirements.skipped:
        _requirements(_sheet(wb, "Đáp ứng yêu cầu"), run)
    _risks(_sheet(wb, "Rủi ro"), run)
    if any(a.data_profiles for a in run.attachments):
        _data(_sheet(wb, "Dữ liệu mẫu"), run)
    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()
