"""Read spreadsheet-like files (.xlsx, .csv) into header + rows of strings/values."""

import csv
import io
from typing import Any

from openpyxl import load_workbook

from app.documents import DocumentError, decode_text

MAX_ROWS = 50_000


class Sheet:
    def __init__(self, name: str | None, rows: list[list[Any]]):
        self.name = name
        self.rows = rows


def read_tables(filename: str, data: bytes) -> list[Sheet]:
    lower = filename.lower()
    if lower.endswith(".csv"):
        text = decode_text(data)
        try:
            dialect = csv.Sniffer().sniff(text[:4096], delimiters=",;\t|")
        except csv.Error:
            dialect = csv.excel
        rows = [
            row
            for _, row in zip(
                range(MAX_ROWS + 50), csv.reader(io.StringIO(text), dialect), strict=False
            )
        ]
        return [Sheet(None, rows)]
    if lower.endswith(".xlsx"):
        try:
            wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        except Exception as exc:  # openpyxl raises many exception types for corrupt files
            raise DocumentError("File Excel bị lỗi hoặc không đúng định dạng .xlsx.") from exc
        sheets = []
        for ws in wb.worksheets:
            rows = []
            for i, row in enumerate(ws.iter_rows(values_only=True)):
                if i >= MAX_ROWS + 50:
                    break
                rows.append(list(row))
            sheets.append(Sheet(ws.title, rows))
        wb.close()
        return sheets
    if lower.endswith(".xls"):
        raise DocumentError("Định dạng .xls cũ chưa được hỗ trợ, hãy lưu lại thành .xlsx.")
    raise DocumentError("Chỉ hỗ trợ bảng tính .xlsx hoặc .csv.")


def cell_text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def trim(rows: list[list[Any]]) -> list[list[Any]]:
    """Drop fully empty rows and trailing empty columns."""
    rows = [r for r in rows if any(cell_text(c) for c in r)]
    width = max(
        (max((i + 1 for i, c in enumerate(r) if cell_text(c)), default=0) for r in rows), default=0
    )
    return [list(r[:width]) + [None] * (width - len(r)) for r in rows]
