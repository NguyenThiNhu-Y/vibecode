"""Builders for simulated attachment files used across tests."""

import io
import zipfile

from docx import Document
from openpyxl import Workbook

REQUIREMENTS = [
    ("Hỏi đáp", "Nhân viên hỏi đáp quy định và nhận câu trả lời kèm trích dẫn", "Must"),
    ("Quản trị", "Phân quyền theo phòng ban, tích hợp SSO hiện có", "Must"),
    ("Mobile", "Ứng dụng di động offline cho nhân viên hiện trường", "Could"),
    ("Khác", "Dùng AI", "Could"),
]


def requirements_xlsx(count: int | None = None) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Yêu cầu"
    ws.append(["DANH SÁCH YÊU CẦU (giả lập)"])
    ws.append([])
    ws.append(["Mã", "Chức năng", "Mô tả yêu cầu", "Ưu tiên"])
    rows = (
        REQUIREMENTS
        if count is None
        else [("Nhóm", f"Yêu cầu số {i} về tra cứu tài liệu", "Must") for i in range(count)]
    )
    for i, (feature, text, priority) in enumerate(rows, start=1):
        ws.append([f"FR-{i:02d}", feature, text, priority])
    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def data_csv(rows: int = 150) -> bytes:
    lines = ["ngay,doanh_thu,email_khach,ghi_chu"]
    lines += [f"2026-01-{i % 28 + 1:02d},{i * 1000},kh{i}@example.vn," for i in range(rows)]
    return "\n".join(lines).encode()


def docx_file(text: str) -> bytes:
    doc = Document()
    doc.add_paragraph(text)
    table = doc.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Hệ thống"
    table.cell(0, 1).text = "SharePoint"
    buffer = io.BytesIO()
    doc.save(buffer)
    return buffer.getvalue()


def code_zip(extra_files: int = 0) -> bytes:
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w") as z:
        z.writestr("proj/app/main.py", "from fastapi import FastAPI\napp = FastAPI()\n")
        z.writestr("proj/requirements.txt", "fastapi\npandas\n")
        z.writestr("proj/web/package.json", '{"dependencies": {"react": "18"}}')
        z.writestr("proj/web/src/App.tsx", "export default 1\n")
        z.writestr("proj/tests/test_main.py", "def test():\n    pass\n")
        z.writestr("proj/node_modules/x/index.js", "ignored\n")
        z.writestr("proj/Dockerfile", "FROM python:3.11\n")
        for i in range(extra_files):
            z.writestr(f"proj/gen/f{i}.txt", "x")
    return buffer.getvalue()
