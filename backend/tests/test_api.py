import io
import json
from collections.abc import Iterator
from datetime import UTC, datetime

import pytest
from fastapi.testclient import TestClient

from app.llm.mock import FIXTURE_DIR, MockLLM
from app.main import app, get_llm_client, get_repo
from app.schemas.run import RunStatus
from app.storage.repo import RunRepo

REQUEST = "Chúng tôi cần trợ lý hỏi đáp quy định nội bộ cho nhân viên (giả lập)."
BLOCKING_GAPS = (FIXTURE_DIR / "scenarios" / "clarify" / "gaps.json").read_text(encoding="utf-8")


@pytest.fixture
def repo(tmp_path) -> RunRepo:
    return RunRepo(tmp_path / "api.db")


@pytest.fixture
def llm_holder() -> dict[str, MockLLM]:
    return {"llm": MockLLM()}


@pytest.fixture
def client(repo: RunRepo, llm_holder: dict[str, MockLLM]) -> Iterator[TestClient]:
    app.dependency_overrides[get_repo] = lambda: repo
    app.dependency_overrides[get_llm_client] = lambda: llm_holder["llm"]
    with TestClient(app) as c:
        yield c
    app.dependency_overrides.clear()


def read_sse(client: TestClient, url: str) -> list[tuple[str, dict]]:
    events: list[tuple[str, dict]] = []
    with client.stream("GET", url) as resp:
        assert resp.status_code == 200
        body = resp.read().decode("utf-8").replace("\r\n", "\n")
    for block in body.split("\n\n"):
        name, data = None, None
        for line in block.splitlines():
            if line.startswith("event:"):
                name = line.split(":", 1)[1].strip()
            elif line.startswith("data:"):
                data = json.loads(line.split(":", 1)[1].strip())
        if name:
            events.append((name, data or {}))
    return events


def create(client: TestClient, text: str = REQUEST) -> str:
    resp = client.post("/api/runs", json={"request_text": text})
    assert resp.status_code == 201
    return resp.json()["run_id"]


def test_health(client: TestClient) -> None:
    assert client.get("/api/health").json() == {"ok": True, "llm_provider": "mock"}


def test_create_run_validates_length(client: TestClient) -> None:
    short = client.post("/api/runs", json={"request_text": "quá ngắn"})
    assert short.status_code == 422
    assert isinstance(short.json()["detail"], str)
    assert client.post("/api/runs", json={"request_text": "x" * 20_001}).status_code == 422
    assert client.post("/api/runs", json={"request_text": " " * 40}).status_code == 422
    run_id = create(client)
    assert len(run_id) == 8 and int(run_id, 16) >= 0


def test_get_run_and_404s(client: TestClient) -> None:
    run_id = create(client)
    body = client.get(f"/api/runs/{run_id}").json()
    assert body["status"] == "created" and body["request_text"] == REQUEST
    assert client.get("/api/runs/missing0").status_code == 404
    assert client.get("/api/runs/missing0").json() == {"detail": "Không tìm thấy run"}
    assert client.get("/api/runs/missing0/stream").status_code == 404


def test_stream_emits_full_sequence(client: TestClient) -> None:
    run_id = create(client)
    events = read_sse(client, f"/api/runs/{run_id}/stream")
    assert [e for e, _ in events] == ["status"] + ["step_started", "step_done"] * 8 + ["status"]
    assert events[0][1] == {"status": "running"} and events[-1][1] == {"status": "done"}
    run = client.get(f"/api/runs/{run_id}").json()
    assert run["status"] == "done" and run["proposal"]["markdown"]


def test_stream_conflict_when_running(client: TestClient, repo: RunRepo) -> None:
    run_id = create(client)
    run = repo.get(run_id)
    assert run is not None
    run.status = RunStatus.RUNNING
    repo.save(run)
    assert client.get(f"/api/runs/{run_id}/stream").status_code == 409


def test_stream_of_finished_run_only_reports_status(client: TestClient) -> None:
    run_id = create(client)
    read_sse(client, f"/api/runs/{run_id}/stream")
    client.post(f"/api/runs/{run_id}/review", json={"approved": True, "note": None})
    assert read_sse(client, f"/api/runs/{run_id}/stream") == [("status", {"status": "approved"})]


def test_clarification_flow(client: TestClient, llm_holder: dict[str, MockLLM]) -> None:
    llm_holder["llm"] = MockLLM({"gaps": [BLOCKING_GAPS]})
    run_id = create(client)
    events = read_sse(client, f"/api/runs/{run_id}/stream")
    assert events[-1] == ("status", {"status": "waiting_clarification"})

    resp = client.post(f"/api/runs/{run_id}/answers", json={"answers": {"q1": "Bộ phận CSKH"}})
    assert resp.status_code == 200
    body = resp.json()
    assert body["status"] == "created" and body["gaps"] is None
    assert body["answers"] == {"q1": "Bộ phận CSKH"}

    llm_holder["llm"] = MockLLM()
    events = read_sse(client, f"/api/runs/{run_id}/stream")
    assert events[-1] == ("status", {"status": "done"})
    assert ("step_started", {"step": "intake"}) not in events


def test_answers_conflict_when_not_waiting(client: TestClient) -> None:
    run_id = create(client)
    resp = client.post(f"/api/runs/{run_id}/answers", json={"answers": {"q1": "x"}})
    assert resp.status_code == 409
    assert client.post("/api/runs/missing0/answers", json={"answers": {}}).status_code == 404


def test_review_flow_and_conflict(client: TestClient) -> None:
    run_id = create(client)
    assert client.post(f"/api/runs/{run_id}/review", json={"approved": True}).status_code == 409
    read_sse(client, f"/api/runs/{run_id}/stream")
    resp = client.post(f"/api/runs/{run_id}/review", json={"approved": False, "note": "Sửa effort"})
    assert resp.status_code == 200
    assert resp.json()["status"] == "rejected" and resp.json()["reviewer_note"] == "Sửa effort"
    again = client.post(f"/api/runs/{run_id}/review", json={"approved": True})
    assert again.status_code == 409


def test_proposal_markdown(client: TestClient) -> None:
    run_id = create(client)
    assert client.get(f"/api/runs/{run_id}/proposal.md").status_code == 409
    read_sse(client, f"/api/runs/{run_id}/stream")
    resp = client.get(f"/api/runs/{run_id}/proposal.md")
    assert resp.status_code == 200
    assert resp.headers["content-type"].startswith("text/markdown")
    assert "## Giả định" in resp.text


def test_list_runs(client: TestClient) -> None:
    first = create(client)
    second = create(client)
    read_sse(client, f"/api/runs/{first}/stream")
    rows = client.get("/api/runs?limit=50").json()
    assert {r["id"] for r in rows} == {first, second}
    done = next(r for r in rows if r["id"] == first)
    assert done["pattern"] == "rag" and done["business_goal"]
    assert {"id", "created_at", "status", "business_goal", "pattern"} <= set(done)
    assert done["deal_stage"] == "reviewing" and done["quote_total"] > 0


def test_failed_step_emits_run_error(client: TestClient, llm_holder: dict[str, MockLLM]) -> None:
    llm_holder["llm"] = MockLLM({"feasibility": ["sai"]})
    run_id = create(client)
    events = read_sse(client, f"/api/runs/{run_id}/stream")
    assert events[-1][0] == "run_error" and events[-1][1]["step"] == "feasibility"
    assert client.get(f"/api/runs/{run_id}").json()["status"] == "failed"


def test_replay_endpoints(client: TestClient, tmp_path, monkeypatch) -> None:
    import app.main as main_module

    recording = {
        "name": "demo",
        "llm_provider": "mock",
        "request_text": "Yêu cầu giả lập",
        "segments": [
            {
                "answers": {},
                "events": [
                    {"event": "status", "data": {"status": "running"}, "delay_ms": 0},
                    {"event": "step_started", "data": {"step": "intake"}, "delay_ms": 5},
                ],
            },
            {
                "answers": {"q1": "Trả lời"},
                "events": [{"event": "status", "data": {"status": "done"}, "delay_ms": 5}],
            },
        ],
    }
    (tmp_path / "demo.json").write_text(json.dumps(recording), encoding="utf-8")
    monkeypatch.setattr(main_module, "REPLAY_DIR", tmp_path)

    assert client.get("/api/replays").json() == ["demo"]
    meta = client.get("/api/replays/demo").json()
    assert meta["segments"][1]["answers"] == {"q1": "Trả lời"}
    events = read_sse(client, "/api/replays/demo/stream?speed=20")
    assert events == [("status", {"status": "running"}), ("step_started", {"step": "intake"})]
    assert read_sse(client, "/api/replays/demo/stream?segment=1") == [
        ("status", {"status": "done"})
    ]
    assert client.get("/api/replays/missing/stream").status_code == 404
    assert client.get("/api/replays/demo/stream?segment=5").status_code == 404
    assert client.get("/api/replays/..%2Fsecret/stream").status_code == 404


def finished_run(client: TestClient) -> str:
    run_id = create(client)
    read_sse(client, f"/api/runs/{run_id}/stream")
    return run_id


def test_rerun_with_feedback(client: TestClient) -> None:
    run_id = create(client)
    body = {"from_step": "architecture", "feedback": "Thêm phương án on-prem"}
    assert client.post(f"/api/runs/{run_id}/rerun", json=body).status_code == 409
    read_sse(client, f"/api/runs/{run_id}/stream")

    assert (
        client.post(f"/api/runs/{run_id}/rerun", json={**body, "feedback": "ngắn"}).status_code
        == 422
    )
    resp = client.post(f"/api/runs/{run_id}/rerun", json=body)
    assert resp.status_code == 200
    run = resp.json()
    assert run["status"] == "created" and run["revision"] == 1
    assert run["architecture"] is None and run["proposal"] is None and run["feasibility"]
    assert run["feedback_step"] == "architecture"

    events = read_sse(client, f"/api/runs/{run_id}/stream")
    started = [d["step"] for e, d in events if e == "step_started"]
    assert started == ["architecture", "wbs", "requirements", "proposal"]
    assert events[-1] == ("status", {"status": "done"})


def test_rerun_allowed_after_rejection_not_after_approval(client: TestClient) -> None:
    run_id = finished_run(client)
    client.post(f"/api/runs/{run_id}/review", json={"approved": False, "note": "Sai effort"})
    body = {"from_step": "pattern", "feedback": "Xem lại hướng giải pháp"}
    assert client.post(f"/api/runs/{run_id}/rerun", json=body).status_code == 200

    other = finished_run(client)
    client.post(f"/api/runs/{other}/review", json={"approved": True})
    assert client.post(f"/api/runs/{other}/rerun", json=body).status_code == 409


def test_edit_proposal(client: TestClient) -> None:
    run_id = create(client)
    edit = {"markdown": "# Bản sửa\n\n## Giả định\n- A\n\n## Rủi ro\n- B", "title": "Bản sửa"}
    assert client.put(f"/api/runs/{run_id}/proposal", json=edit).status_code == 409
    read_sse(client, f"/api/runs/{run_id}/stream")
    resp = client.put(f"/api/runs/{run_id}/proposal", json=edit)
    assert resp.status_code == 200
    run = resp.json()
    assert run["proposal"]["title"] == "Bản sửa" and run["proposal_edited"] is True
    assert client.get(f"/api/runs/{run_id}/proposal.md").text == edit["markdown"]
    client.post(f"/api/runs/{run_id}/review", json={"approved": True})
    assert client.put(f"/api/runs/{run_id}/proposal", json=edit).status_code == 409


def test_translate_proposal_is_cached(client: TestClient, llm_holder: dict[str, MockLLM]) -> None:
    run_id = finished_run(client)
    llm = MockLLM()
    llm_holder["llm"] = llm
    resp = client.post(f"/api/runs/{run_id}/proposal/translate", json={"language": "en"})
    assert resp.status_code == 200 and "Preliminary proposal" in resp.json()["markdown"]
    client.post(f"/api/runs/{run_id}/proposal/translate", json={"language": "en"})
    assert llm.calls == ["translate"]
    same = client.post(f"/api/runs/{run_id}/proposal/translate", json={"language": "vi"})
    assert same.json()["markdown"].startswith("# Đề xuất sơ bộ")
    assert client.get(f"/api/runs/{run_id}").json()["translations"].keys() == {"en"}
    bad = client.post(f"/api/runs/{run_id}/proposal/translate", json={"language": "fr"})
    assert bad.status_code == 422


def test_translate_failure_returns_502(client: TestClient, llm_holder: dict[str, MockLLM]) -> None:
    run_id = finished_run(client)
    llm_holder["llm"] = MockLLM({"translate": ["không phải json"]})
    resp = client.post(f"/api/runs/{run_id}/proposal/translate", json={"language": "ja"})
    assert resp.status_code == 502 and "Không dịch được" in resp.json()["detail"]


def test_extract_document_endpoint(client: TestClient) -> None:
    from tests.test_documents import PDF_TEXT, make_pdf

    pdf = client.post(
        "/api/documents/extract", files={"file": ("rfp.pdf", make_pdf(PDF_TEXT), "application/pdf")}
    )
    assert pdf.status_code == 200 and "maintenance manuals" in pdf.json()["text"]
    bad = client.post(
        "/api/documents/extract", files={"file": ("a.exe", b"x", "application/octet-stream")}
    )
    assert bad.status_code == 422 and "Chỉ hỗ trợ" in bad.json()["detail"]


def test_eval_report_endpoints(client: TestClient, tmp_path, monkeypatch) -> None:
    import app.main as main_module

    report = {
        "provider": "mock",
        "label": "v1",
        "created_at": "2026-10-02T10:00:00",
        "summary": {"pattern_accuracy": 1.0},
        "cases": [{"id": "case_01", "outputs": {"big": True}}],
    }
    (tmp_path / "eval_1002_1000.json").write_text(json.dumps(report), encoding="utf-8")
    (tmp_path / "eval_broken.json").write_text("{", encoding="utf-8")
    monkeypatch.setattr(main_module, "EVAL_REPORT_DIR", tmp_path)

    listing = client.get("/api/eval/reports").json()
    assert [r["name"] for r in listing] == ["eval_1002_1000"]
    assert listing[0]["summary"] == {"pattern_accuracy": 1.0} and listing[0]["label"] == "v1"
    detail = client.get("/api/eval/reports/eval_1002_1000").json()
    assert detail["cases"] == [{"id": "case_01"}]
    assert client.get("/api/eval/reports/missing").status_code == 404


def test_startup_resets_runs_stuck_in_running(repo: RunRepo, llm_holder) -> None:
    from app.schemas.run import ScopingRun

    stuck = ScopingRun(
        id="stuck001", created_at=datetime.now(UTC), request_text="x" * 40, status=RunStatus.RUNNING
    )
    repo.save(stuck)
    app.dependency_overrides[get_repo] = lambda: repo
    try:
        with TestClient(app):
            pass
    finally:
        app.dependency_overrides.clear()
    stored = repo.get("stuck001")
    assert stored is not None and stored.status == RunStatus.CREATED


def test_upload_attachments_and_exports(client: TestClient) -> None:
    from tests.samples import code_zip, data_csv, requirements_xlsx

    run_id = create(client)
    files = [
        ("files", ("req.xlsx", requirements_xlsx(), "application/octet-stream")),
        ("files", ("data.csv", data_csv(), "text/csv")),
        ("files", ("code.zip", code_zip(), "application/zip")),
    ]
    resp = client.post(
        f"/api/runs/{run_id}/attachments", files=files, data={"kinds": ["auto", "", "source_code"]}
    )
    assert resp.status_code == 200
    assert [a["kind"] for a in resp.json()["attachments"]] == [
        "requirements",
        "data_sample",
        "source_code",
    ]

    removed = client.delete(f"/api/runs/{run_id}/attachments/a2").json()
    assert [a["id"] for a in removed["attachments"]] == ["a1", "a3"]
    assert client.delete(f"/api/runs/{run_id}/attachments/zz").status_code == 404

    assert client.get(f"/api/runs/{run_id}/export/slides.pptx").status_code == 409
    read_sse(client, f"/api/runs/{run_id}/stream")
    for name, media, doc in [
        ("slides.pptx", "presentationml", "Slides"),
        ("proposal.docx", "wordprocessingml", "Proposal"),
        ("workbook.xlsx", "spreadsheetml", "Estimate"),
        ("package.zip", "application/zip", "Package"),
    ]:
        export = client.get(f"/api/runs/{run_id}/export/{name}")
        assert export.status_code == 200 and media in export.headers["content-type"]
        # company naming rule: {company}_{client}_{project}_{doc}_v{version}_{date}
        assert f"ABC_KH_{run_id}_{doc}_v0.1_" in export.headers["content-disposition"]
    assert client.get(f"/api/runs/{run_id}/export/evil.exe").status_code == 404

    late = client.post(f"/api/runs/{run_id}/attachments", files=[files[0]])
    assert late.status_code == 409


def test_upload_attachment_errors(client: TestClient) -> None:
    run_id = create(client)
    bad = client.post(
        f"/api/runs/{run_id}/attachments",
        files=[("files", ("x.exe", b"x", "application/octet-stream"))],
    )
    assert bad.status_code == 422 and "x.exe" in bad.json()["detail"]
    wrong_kind = client.post(
        f"/api/runs/{run_id}/attachments",
        files=[("files", ("a.txt", "Nội dung đủ dài để đọc được".encode(), "text/plain"))],
        data={"kinds": ["unknown"]},
    )
    assert wrong_kind.status_code == 422
    many = [("files", (f"f{i}.txt", "Nội dung tài liệu".encode(), "text/plain")) for i in range(11)]
    assert client.post(f"/api/runs/{run_id}/attachments", files=many).status_code == 422


def test_create_with_meta_and_patch(client: TestClient) -> None:
    resp = client.post(
        "/api/runs",
        json={
            "request_text": REQUEST,
            "project_name": "Trợ lý quy định",
            "client_name": "Ngân hàng A (giả lập)",
            "due_date": "2026-10-20",
        },
    )
    run_id = resp.json()["run_id"]
    run = client.get(f"/api/runs/{run_id}").json()
    assert run["project_name"] == "Trợ lý quy định" and run["due_date"] == "2026-10-20"
    assert run["deal_stage"] == "new"
    patched = client.patch(
        f"/api/runs/{run_id}/meta", json={"deal_stage": "won", "client_name": "  "}
    ).json()
    assert patched["deal_stage"] == "won" and patched["client_name"] is None
    assert patched["project_name"] == "Trợ lý quy định"
    assert client.patch(f"/api/runs/{run_id}/meta", json={"deal_stage": "magic"}).status_code == 422
    row = client.get("/api/runs").json()[0]
    assert row["project_name"] == "Trợ lý quy định" and row["deal_stage"] == "won"


def test_quotation_endpoint(client: TestClient) -> None:
    run_id = create(client)
    assert client.post(f"/api/runs/{run_id}/quotation", json={}).status_code == 409
    read_sse(client, f"/api/runs/{run_id}/stream")
    run = client.get(f"/api/runs/{run_id}").json()
    assert run["quotation"]["currency"] == "VND" and run["deal_stage"] == "reviewing"
    usd = client.post(
        f"/api/runs/{run_id}/quotation", json={"currency": "USD", "contingency_pct": 20}
    ).json()
    assert usd["quotation"]["currency"] == "USD" and usd["quotation"]["contingency_pct"] == 20
    again = client.post(f"/api/runs/{run_id}/quotation", json={}).json()
    assert again["quotation"]["currency"] == "USD"  # keeps previous choices
    assert client.post(f"/api/runs/{run_id}/quotation", json={"currency": "EUR"}).status_code == 422
    tm = client.post(
        f"/api/runs/{run_id}/quotation",
        json={"contract_model": "time_material", "onsite_ratio": 20},
    ).json()["quotation"]
    assert tm["contract_model"] == "time_material" and tm["contingency_pct"] == 0
    assert tm["onsite_ratio"] == 20 and tm["currency"] == "USD"
    # technical approval alone is not enough: the price needs its own sign-off
    client.post(f"/api/runs/{run_id}/review", json={"approved": True})
    assert client.get(f"/api/runs/{run_id}").json()["deal_stage"] == "reviewing"
    url = f"/api/runs/{run_id}/pricing-approval"
    assert client.post(url, json={"approved": False}).status_code == 422  # reason required
    assert client.post(url, json={"approved": True}).json()["deal_stage"] == "ready"
    assert client.post(f"/api/runs/{run_id}/quotation", json={}).status_code == 409  # locked
    client.post(url, json={"approved": False, "note": "Giảm onsite xuống 10%"})
    again = client.post(f"/api/runs/{run_id}/quotation", json={"onsite_ratio": 10}).json()
    assert again["pricing_approval"] is None and again["deal_stage"] == "reviewing"


def test_client_email_fills_names_by_code(
    client: TestClient, llm_holder: dict[str, MockLLM]
) -> None:
    resp = client.post(
        "/api/runs", json={"request_text": REQUEST, "client_name": "Công ty B (giả lập)"}
    )
    run_id = resp.json()["run_id"]
    assert client.post(f"/api/runs/{run_id}/client-email", json={}).status_code == 409
    read_sse(client, f"/api/runs/{run_id}/stream")
    llm = MockLLM()
    llm_holder["llm"] = llm
    email = client.post(
        f"/api/runs/{run_id}/client-email", json={"sender_name": "Nguyễn Văn A"}
    ).json()
    assert email["body"].startswith("Kính gửi Công ty B (giả lập),")
    assert "Nguyễn Văn A" in email["body"] and "{{" not in email["body"]
    assert "1. " in email["body"]
    sent_to_llm = " ".join(llm.calls)
    assert sent_to_llm == "client_email"
    client.post(f"/api/runs/{run_id}/client-email", json={})
    assert llm.calls == ["client_email"]  # cached
    client.post(f"/api/runs/{run_id}/client-email", json={"regenerate": True})
    assert llm.calls == ["client_email", "client_email"]


def test_qa_sheet_export_and_import(client: TestClient, llm_holder: dict[str, MockLLM]) -> None:
    from openpyxl import load_workbook

    llm_holder["llm"] = MockLLM({"gaps": [BLOCKING_GAPS]})
    run_id = create(client)
    read_sse(client, f"/api/runs/{run_id}/stream")
    sheet = client.get(f"/api/runs/{run_id}/export/qa_sheet.xlsx")
    assert sheet.status_code == 200
    wb = load_workbook(io.BytesIO(sheet.content))
    ws = wb.active
    last = ws.max_column
    ws.cell(row=5, column=last, value="Bộ phận chăm sóc khách hàng")
    ws.cell(row=6, column=last, value="Email và FAQ")
    buffer = io.BytesIO()
    wb.save(buffer)
    files = {"file": ("qa.xlsx", buffer.getvalue(), "application/octet-stream")}
    resp = client.post(f"/api/runs/{run_id}/answers/import", files=files)
    assert resp.status_code == 200
    assert resp.json()["answers"] == {"q1": "Bộ phận chăm sóc khách hàng", "q2": "Email và FAQ"}
    assert resp.json()["status"] == "created"
    assert client.post(f"/api/runs/{run_id}/answers/import", files=files).status_code == 409


def test_similar_and_stats(client: TestClient) -> None:
    first = finished_run(client)
    second = finished_run(client)
    similar = client.get(f"/api/runs/{first}/similar").json()
    assert similar and similar[0]["id"] == second and similar[0]["score"] >= 0.25
    data = client.get("/api/stats").json()
    assert data["runs"] == 2 and data["finished"] == 2 and data["stages"]["reviewing"] == 2
    assert data["hours_saved"] > 40 and data["win_rate"] is None


def test_localized_slides_export(client: TestClient) -> None:
    run_id = finished_run(client)
    ja = client.get(f"/api/runs/{run_id}/export/slides.pptx?lang=ja")
    assert ja.status_code == 200 and "_Slides-JA_" in ja.headers["content-disposition"]
    assert client.get(f"/api/runs/{run_id}").json()["deck_translations"]["ja"]
    assert client.get(f"/api/runs/{run_id}/export/slides.pptx?lang=fr").status_code == 422


def test_settings_endpoints(client: TestClient, tmp_path, monkeypatch) -> None:
    import shutil

    from app.knowledge import loader

    target = tmp_path / "kb"
    shutil.copytree(loader.KB_DIR, target)
    monkeypatch.setattr(loader, "KB_DIR", target)
    loader.reload()
    try:
        data = client.get("/api/settings").json()
        assert {"rate_card", "estimation_template", "reference_projects"} <= set(data)
        card = data["rate_card"]
        card["contingency"]["base_pct"] = 12
        assert (
            client.put("/api/settings/rate-card", json=card).json()["contingency"]["base_pct"] == 12
        )
        card["default_role"] = "unknown"
        bad = client.put("/api/settings/rate-card", json=card)
        assert bad.status_code == 422 and "default_role" in bad.json()["detail"]
        assert (
            client.put(
                "/api/settings/reference-projects/rp_new_one",
                json={"content": "Nội dung dự án giả lập đủ dài"},
            ).status_code
            == 200
        )
        assert client.delete("/api/settings/reference-projects/rp_new_one").status_code == 200
        assert client.delete("/api/settings/reference-projects/rp_new_one").status_code == 404
    finally:
        loader.reload()


def test_replay_export_from_snapshot(
    client: TestClient, tmp_path, monkeypatch, repo: RunRepo
) -> None:
    import app.main as main_module

    run_id = finished_run(client)
    snapshot = client.get(f"/api/runs/{run_id}").json()
    recording = {
        "name": "demo",
        "segments": [{"answers": {}, "events": []}],
        "final_run": snapshot,
        "attachments": [],
    }
    (tmp_path / "demo.json").write_text(json.dumps(recording), encoding="utf-8")
    old = {**recording, "name": "old", "final_run": None}
    (tmp_path / "old.json").write_text(json.dumps(old), encoding="utf-8")
    monkeypatch.setattr(main_module, "REPLAY_DIR", tmp_path)
    assert client.get("/api/replays/demo").json()["has_final_run"] is True
    resp = client.get("/api/replays/demo/export/workbook.xlsx")
    assert resp.status_code == 200 and "spreadsheetml" in resp.headers["content-type"]
    assert client.get("/api/replays/old").json()["has_final_run"] is False
    assert client.get("/api/replays/old/export/workbook.xlsx").status_code == 409


def test_bid_versions_and_case_studies(client: TestClient) -> None:
    run_id = finished_run(client)
    bid = client.get(f"/api/runs/{run_id}/bid").json()
    assert bid["recommendation"] == "need_info" and bid["decision"] is None
    manual = {c["id"]: True for c in bid["criteria"] if not c["auto"]}
    assert (
        client.put(f"/api/runs/{run_id}/bid", json={"checks": {"magic": True}}).status_code == 422
    )
    decided = client.put(f"/api/runs/{run_id}/bid", json={"checks": manual, "decision": "no_bid"})
    assert decided.json()["decision"] == "no_bid" and decided.json()["recommendation"] == "bid"
    assert client.get(f"/api/runs/{run_id}").json()["deal_stage"] == "no_bid"
    client.put(f"/api/runs/{run_id}/bid", json={"checks": manual, "decision": "bid"})
    assert client.get(f"/api/runs/{run_id}").json()["deal_stage"] == "reviewing"

    url = f"/api/runs/{run_id}/versions"
    assert client.post(url, json={"note": "Bản đầu"}).json()["versions"][-1]["version"] == "1.0"
    assert client.post(url, json={}).json()["versions"][-1]["version"] == "1.1"
    major = client.post(url, json={"major": True, "sent": True}).json()
    assert major["versions"][-1]["version"] == "2.0" and major["deal_stage"] == "sent"
    assert major["versions"][-1]["markdown"] and major["versions"][-1]["total"]
    sent = client.patch(f"{url}/1.0", json={"sent": True}).json()
    assert sent["versions"][0]["sent"] and sent["versions"][0]["sent_at"]
    assert client.patch(f"{url}/9.9", json={"sent": True}).status_code == 404
    export = client.get(f"/api/runs/{run_id}/export/proposal.docx")
    assert "_Proposal_v2.0_" in export.headers["content-disposition"]

    cases = client.get(f"/api/runs/{run_id}/case-studies").json()
    assert cases and cases[0]["pattern"] == "rag" and cases[0]["reasons"]


def test_company_settings_and_templates(client: TestClient, tmp_path, monkeypatch) -> None:
    import shutil

    from pptx import Presentation

    from app.knowledge import loader

    target = tmp_path / "kb"
    shutil.copytree(loader.KB_DIR, target)
    monkeypatch.setattr(loader, "KB_DIR", target)
    data = client.get("/api/settings").json()
    assert {"company", "content_library", "case_studies", "bid_criteria"} <= set(data)
    company = {**data["company"], "short_name": "XYZ"}
    assert client.put("/api/settings/company", json=company).json()["short_name"] == "XYZ"
    bad = client.put("/api/settings/company", json={**company, "file_naming": "{company}"})
    assert bad.status_code == 422 and "{doc}" in bad.json()["detail"]
    blocks = data["content_library"]
    assert client.put("/api/settings/content-library", json=blocks + blocks[:1]).status_code == 422
    assert (
        client.put("/api/settings/case-studies", json=data["case_studies"][:2]).status_code == 200
    )
    assert client.put("/api/settings/bid-criteria", json=[]).status_code == 422

    listing = client.get("/api/templates").json()
    assert [i["kind"] for i in listing["items"]] == [
        "slides",
        "proposal_docx",
        "workbook",
        "qa_sheet",
    ]
    sample = client.get("/api/templates/slides/sample/file")
    assert sample.status_code == 200 and "presentationml" in sample.headers["content-type"]
    assert client.get("/api/templates/slides/custom/file").status_code == 404
    assert (
        client.post("/api/templates/slides", files={"file": ("x.pptx", b"nope")}).status_code == 422
    )
    plain = io.BytesIO()
    Presentation().save(plain)
    up = client.post("/api/templates/slides", files={"file": ("deck.pptx", plain.getvalue())})
    assert up.status_code == 200 and up.json()["warnings"]
    assert up.json()["items"][0]["active"] == "custom" and up.json()["items"][0]["custom"]
    active = client.put("/api/templates/active", json={"kind": "slides", "source": "sample"})
    assert active.json()["items"][0]["active"] == "sample"
    cfg = listing["config"]
    cfg["qa_sheet"]["columns"]["answer"] = "G"
    assert (
        client.put("/api/templates/config", json=cfg).json()["config"]["qa_sheet"]["columns"][
            "answer"
        ]
        == "G"
    )
    cfg["qa_sheet"]["columns"]["answer"] = "g1"
    assert client.put("/api/templates/config", json=cfg).status_code == 422
    assert client.delete("/api/templates/slides/custom").status_code == 200
    assert client.delete("/api/templates/slides/custom").status_code == 404


def test_logo_endpoints(client: TestClient, tmp_path, monkeypatch) -> None:
    import shutil

    from PIL import Image

    from app.knowledge import loader
    from app.templates_store import logo_path

    target = tmp_path / "kb"
    shutil.copytree(loader.KB_DIR, target)
    monkeypatch.setattr(loader, "KB_DIR", target)
    logo_path(target).unlink(missing_ok=True)
    assert client.get("/api/settings").json()["logo"] is None
    assert client.get("/api/settings/logo").status_code == 404
    bad = client.post("/api/settings/logo", files={"file": ("logo.png", b"nope")})
    assert bad.status_code == 422 and "PNG" in bad.json()["detail"]
    png = io.BytesIO()
    Image.new("RGBA", (300, 100), (0, 0, 0, 255)).save(png, format="PNG")
    up = client.post("/api/settings/logo", files={"file": ("logo.png", png.getvalue())})
    assert up.status_code == 200 and up.json()["logo"]["width"] == 300
    assert client.get("/api/settings").json()["logo"]["height"] == 100
    got = client.get("/api/settings/logo")
    assert got.status_code == 200 and got.headers["content-type"] == "image/png"
    assert client.delete("/api/settings/logo").json() == {"logo": None}
    assert client.delete("/api/settings/logo").status_code == 404


def test_bidding_workbook_endpoint(client: TestClient) -> None:
    from openpyxl import load_workbook

    run_id = create(client)
    assert client.get(f"/api/runs/{run_id}/bidding.xlsx").status_code == 409  # no gaps yet
    waiting = create(client, "Chúng tôi muốn dùng AI để tăng năng suất cho nhân viên văn phòng.")
    read_sse(client, f"/api/runs/{waiting}/stream")
    early = client.get(f"/api/runs/{waiting}/bidding.xlsx")
    assert early.status_code == 200 and "attachment" in early.headers["content-disposition"]
    assert load_workbook(io.BytesIO(early.content)).sheetnames == ["Q&A"]

    read_sse(client, f"/api/runs/{run_id}/stream")
    full = client.get(f"/api/runs/{run_id}/bidding.xlsx")
    assert full.headers["content-type"].startswith("application/vnd.openxmlformats")
    assert "Bidding" in full.headers["content-disposition"]
    names = load_workbook(io.BytesIO(full.content)).sheetnames
    assert names == ["Q&A", "WBS", "Summary", "Master Schedule"]
    same = client.get(f"/api/runs/{run_id}/export/bidding.xlsx")
    assert same.status_code == 200


def test_schedule_config_recomputes_without_llm(client: TestClient) -> None:
    run_id = create(client)
    config = {"start_date": "2026-11-02", "headcount": {"AI": 2, "BE": 1, "FE": 1, "QA": 1,
              "BA": 1, "PM": 1, "INFRA": 1, "DESIGN": 1, "DATA": 1}}  # fmt: skip
    url = f"/api/runs/{run_id}/schedule-config"
    assert client.put(url, json=config).status_code == 409  # no WBS yet
    read_sse(client, f"/api/runs/{run_id}/stream")

    two = client.put(url, json=config).json()
    assert two["schedule"]["phases"][0]["start"] == "2026-11-02"
    assert two["schedule"]["milestones"][0] == {
        "id": "M1", "name": "Kick-off", "date": "2026-11-02", "phase": "poc", "payment_percent": 30,
    }  # fmt: skip
    four = client.put(url, json={**config, "headcount": {**config["headcount"], "AI": 4}}).json()
    mvp = lambda run: next(p for p in run["schedule"]["phases"] if p["phase"] == "mvp")  # noqa: E731
    assert mvp(four)["end"] <= mvp(two)["end"]
    assert four["schedule_config"]["headcount"]["AI"] == 4
    assert four["quotation"]["months"] is not None  # quotation follows the new timeline

    zero = client.put(url, json={**config, "headcount": {**config["headcount"], "QA": 0}})
    assert zero.status_code == 422 and "QA" in zero.json()["detail"]
    assert client.get(f"/api/runs/{run_id}").json()["schedule_config"]["headcount"]["AI"] == 4


def test_create_run_with_schedule_config(client: TestClient) -> None:
    headcount = {"AI": 3, "BE": 2, "FE": 1, "QA": 1, "BA": 1, "PM": 1, "INFRA": 1, "DESIGN": 1}
    config = {"start_date": "2026-12-07", "headcount": headcount | {"DATA": 1}, "buffer_ratio": 0.1}
    resp = client.post("/api/runs", json={"request_text": REQUEST, "schedule_config": config})
    run_id = resp.json()["run_id"]
    read_sse(client, f"/api/runs/{run_id}/stream")
    run = client.get(f"/api/runs/{run_id}").json()
    assert run["schedule"]["phases"][0]["start"] == "2026-12-07"
    assert run["schedule"]["config"]["buffer_ratio"] == 0.1


def test_edit_wbs_estimates_recomputes_by_code(client: TestClient) -> None:
    run_id = create(client)
    url = f"/api/runs/{run_id}/wbs"
    assert client.patch(url, json={"estimates": {"1.1": 2}}).status_code == 409  # no WBS yet
    read_sse(client, f"/api/runs/{run_id}/stream")
    before = client.get(f"/api/runs/{run_id}").json()
    items = before["wbs"]["items"]
    parents = {i["id"].rsplit(".", 1)[0] for i in items if "." in i["id"]}
    leaf = next(i for i in items if i["id"] not in parents and i["phase"] == "mvp")
    parent = leaf["id"].rsplit(".", 1)[0]
    old_parent = next(i["estimate_md"] for i in items if i["id"] == parent)
    new_value = 10 if leaf["estimate_md"] != 10 else 1

    after = client.patch(url, json={"estimates": {leaf["id"]: new_value}}).json()
    delta = new_value - leaf["estimate_md"]
    by_id = {i["id"]: i for i in after["wbs"]["items"]}
    assert by_id[leaf["id"]]["estimate_md"] == new_value and after["wbs_edited"] is True
    assert by_id[parent]["estimate_md"] == old_parent + delta  # parent re-summed by code
    mvp = lambda run: next(t for t in run["wbs"]["totals"] if t["phase"] == "mvp")  # noqa: E731
    assert mvp(after)["total_md"] == mvp(before)["total_md"] + delta
    assert after["quotation"]["wbs_person_days"] == before["quotation"]["wbs_person_days"] + delta
    assert after["schedule"]["config"] == before["schedule"]["config"]  # same start / headcount

    for body, fragment in [
        ({"estimates": {parent: 5}}, "node cha"),
        ({"estimates": {"9.9.9": 1}}, "Không có task"),
        ({"estimates": {leaf["id"]: 11}}, ""),
        ({"estimates": {}}, ""),
    ]:
        resp = client.patch(url, json=body)
        assert resp.status_code == 422 and fragment in str(resp.json()["detail"])

    approve = {"approved": True, "note": None}
    assert client.post(f"/api/runs/{run_id}/pricing-approval", json=approve).status_code == 200
    locked = client.patch(url, json={"estimates": {leaf["id"]: 3}})
    assert locked.status_code == 409 and "duyệt giá" in locked.json()["detail"]

    body = {"from_step": "wbs", "feedback": "Tách nhỏ task tích hợp hơn"}
    client.post(f"/api/runs/{run_id}/pricing-approval", json={"approved": False, "note": "sửa"})
    assert client.post(f"/api/runs/{run_id}/rerun", json=body).json()["wbs_edited"] is False
