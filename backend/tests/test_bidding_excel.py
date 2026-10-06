"""bidding.xlsx (docs/BIDDING_SPEC.md 6.1, task B06) and question priority (3.1)."""

import io
import re

from openpyxl import load_workbook

from app.agents.pipeline import run_pipeline
from app.agents.steps import validate_gaps
from app.exports.bidding import build_bidding_xlsx
from app.exports.plan import week_starts
from app.llm.mock import MockLLM
from app.schemas.gaps import GapResult
from app.schemas.run import RunStatus
from tests.conftest import MemoryRepo, load_fixture, make_run
from tests.test_exports import finished_run

SHEETS = ["Q&A", "WBS", "Summary", "Master Schedule"]


def headers(ws) -> list:
    return [c.value for c in ws[1] if c.value is not None]


def sum_value(ws, ref: str, seen: frozenset = frozenset()) -> float:
    """Evaluate the =SUM(Gx,Gy,...) chains the WBS sheet uses (openpyxl does not compute)."""
    value = ws[ref].value
    if isinstance(value, str) and value.startswith("=SUM("):
        refs = re.findall(r"G\d+", value)
        assert ref not in seen and refs, f"{ref}: {value}"
        return sum(sum_value(ws, r, seen | {ref}) for r in refs)
    return float(value or 0)


async def test_full_workbook_sheets_headers_and_formulas() -> None:
    run = await finished_run()
    wb = load_workbook(io.BytesIO(build_bidding_xlsx(run)))
    assert wb.sheetnames == SHEETS
    qa, wbs, summary, master = (wb[n] for n in SHEETS)
    assert headers(qa) == ["No", "Category", "Question", "Why we ask", "Priority", "Blocking",
                           "Customer answer", "Status"]  # fmt: skip
    assert headers(wbs) == ["ID", "Phase", "Task", "Sub-task", "Type", "Priority",
                            "Estimate (MD)", "Depends on", "Deliverable", "Note"]  # fmt: skip
    for ws in (qa, wbs, summary, master):
        assert ws.freeze_panes in ("A2", "G2") and ws.auto_filter.ref
        assert ws["A1"].font.bold and ws["A1"].fill.fgColor.rgb.endswith("D9D9D9")

    rows = {wbs.cell(row=r, column=1).value: r for r in range(2, wbs.max_row + 1)}
    for item in run.wbs.items:
        cell = wbs[f"G{rows[item.id]}"]
        children = [
            c.id for c in run.wbs.items if c.id.rsplit(".", 1)[0] == item.id and "." in c.id
        ]
        if children:  # parent = formula over exactly its direct children
            assert cell.value.startswith("=SUM(")
            referenced = {wbs[f"A{r}"].value for r in re.findall(r"G(\d+)", cell.value)}
            assert referenced == set(children)
        else:
            assert cell.value == item.estimate_md
    grand = next(r for r in range(2, wbs.max_row + 1) if wbs[f"C{r}"].value == "Grand total")
    assert sum_value(wbs, f"G{grand}") == sum(t.total_md for t in run.wbs.totals)
    level3 = next(i for i in run.wbs.items if i.level == 3)
    assert wbs[f"D{rows[level3.id]}"].value == level3.name  # sub-task column

    assert summary["B2"].value.startswith("=SUMIFS(WBS!$G:$G,WBS!$B:$B,$A2,WBS!$E:$E,B$1)")
    assert summary["A2"].value == "PoC" and headers(summary)[-1] == "Total"
    assert any(summary.cell(row=r, column=1).value == "Assumptions" for r in range(1, 40))

    weeks = week_starts(run.schedule)
    assert len(headers(master)) == 6 + len(weeks) and master["G1"].value == f"{weeks[0]:%d/%m}"
    ids = [master.cell(row=r, column=1).value for r in range(2, master.max_row + 1)]
    assert {t.wbs_id for t in run.schedule.tasks} <= set(ids) and "M4" in ids
    m4 = ids.index("M4") + 2
    assert "◆" in [master.cell(row=m4, column=c).value for c in range(7, 7 + len(weeks))]


async def test_qa_sheet_only_while_waiting_for_answers() -> None:
    run = make_run("Chúng tôi muốn dùng AI để tăng năng suất cho nhân viên văn phòng.")
    async for _ in run_pipeline(MockLLM(), run, MemoryRepo()):
        pass
    assert run.status == RunStatus.WAITING_CLARIFICATION and run.wbs is None
    wb = load_workbook(io.BytesIO(build_bidding_xlsx(run)))
    assert wb.sheetnames == ["Q&A"]
    qa = wb["Q&A"]
    assert qa.max_row == len(run.gaps.questions) + 1
    high = [r for r in range(2, qa.max_row + 1) if qa[f"E{r}"].value == "High"]
    assert high and qa[f"A{high[0]}"].fill.fgColor.rgb.endswith("FDE2E1")  # light red
    assert {qa[f"H{r}"].value for r in range(2, qa.max_row + 1)} == {"Open"}


def test_blocking_question_is_always_high_priority() -> None:
    data = load_fixture("gaps")
    data["questions"][0] |= {"blocking": True, "priority": "low"}
    result = validate_gaps(GapResult.model_validate(data))
    assert result.questions[0].priority.value == "high" and not result.can_proceed
