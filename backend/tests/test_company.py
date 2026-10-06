"""Company standards: content library, case studies, bid/no-bid, naming, company templates."""

import io
import re
import shutil
import zipfile

import pytest
from docx import Document
from openpyxl import load_workbook
from PIL import Image
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

from app.agents.pipeline import run_pipeline
from app.bid import evaluate_bid, load_bid_criteria
from app.company import (
    content_disposition,
    export_filename,
    load_case_studies,
    load_company,
    load_content_blocks,
    match_case_studies,
    render_blocks,
)
from app.exports.document import build_docx
from app.exports.qa_sheet import build_qa_sheet, parse_qa_answers
from app.exports.slides import build_slides
from app.exports.workbook import build_workbook
from app.knowledge import loader
from app.llm.mock import MockLLM
from app.schemas.deal import BidDecision, ProposalVersion
from app.templates_store import (
    ExportKit,
    TemplateError,
    delete_custom,
    delete_logo,
    load_templates_config,
    logo_path,
    save_custom,
    save_logo,
    set_active,
    template_path,
)
from tests.conftest import MemoryRepo, make_run


async def finished(text: str | None = None):
    run = make_run(text) if text else make_run()
    async for _ in run_pipeline(MockLLM(), run, MemoryRepo()):
        pass
    return run


@pytest.fixture
def kb(tmp_path, monkeypatch):
    target = tmp_path / "kb"
    shutil.copytree(loader.KB_DIR, target)
    monkeypatch.setattr(loader, "KB_DIR", target)
    logo_path(target).unlink(missing_ok=True)  # a logo uploaded on this machine must not leak in
    return target


def image_bytes(size: tuple[int, int], fmt: str = "PNG") -> bytes:
    buffer = io.BytesIO()
    Image.new("RGB" if fmt == "JPEG" else "RGBA", size, (20, 60, 160)).save(buffer, format=fmt)
    return buffer.getvalue()


# ---------- content library / naming ----------
async def test_blocks_fill_placeholders_in_customer_language() -> None:
    run = await finished()
    company = load_company()
    blocks = load_content_blocks()
    start = render_blocks(run, company, blocks, "start", "ja")
    assert start and start[0][0] == f"{company.short_name}について"
    assert company.name in start[0][1] and "{{" not in start[0][1]
    assert len(render_blocks(run, company, blocks, "end", "vi")) == 4


def test_file_name_follows_company_rule() -> None:
    run = make_run()
    run.client_name, run.project_name = "Ngân hàng Đông Á/XYZ", "Trợ lý quy định"
    run.versions = [ProposalVersion(version="1.1", created_at=run.created_at)]
    name = export_filename(run, load_company(), "proposal", "docx")
    assert re.fullmatch(
        r"ABC_Ngân-hàng-Đông-Á-XYZ_Trợ-lý-quy-định_Proposal_v1\.1_\d{8}\.docx", name
    )
    header = content_disposition(name)
    assert (
        'filename="ABC_Ngan-hang-Dong-A-XYZ_' in header
        and "filename*=UTF-8''ABC_Ng%C3%A2n" in header
    )


# ---------- case studies ----------
async def test_case_studies_match_pattern_and_skip_private() -> None:
    run = await finished()  # mock scenario: RAG, Vietnamese bank
    studies = load_case_studies()
    matches = match_case_studies(run, studies)
    assert matches and all(m["case"].pattern == "rag" for m in matches)
    assert matches[0]["case"].market == "vn" and "Cùng hướng giải pháp" in matches[0]["reasons"]
    for cs in studies:
        cs.public = False
    assert match_case_studies(run, studies) == []


# ---------- bid / no-bid ----------
async def test_bid_suggestions_and_human_override() -> None:
    run = await finished()
    criteria = load_bid_criteria()
    result = evaluate_bid(run, criteria)
    rows = {r["id"]: r for r in result["criteria"]}
    assert rows["solution_clear"]["suggested"] is True and rows["relationship"]["suggested"] is None
    assert rows["budget_known"]["suggested"] is False  # the sample request has no budget
    assert result["recommendation"] == "need_info"  # manual criteria not answered yet
    run.bid = BidDecision(checks={c.id: True for c in criteria if c.auto is None})
    assert evaluate_bid(run, criteria)["recommendation"] == "bid"
    run.bid.checks.update({"capacity": False, "margin_ok": False, "relationship": False})
    assert evaluate_bid(run, criteria)["recommendation"] in {"consider", "no_bid"}


# ---------- company templates ----------
async def test_slides_follow_sample_template() -> None:
    run = await finished()
    run.client_name = "Khách hàng mẫu"
    data = build_slides(run, kit=ExportKit.load(source="sample"))
    prs = Presentation(io.BytesIO(data))
    assert len(prs.slides) == 15  # no attachments: no inputs / requirements slides
    assert prs.slides[0].slide_layout.name == "Title Slide"
    assert all(s.slide_layout.name == "Title Only" for s in list(prs.slides)[1:])
    xml = "".join(
        part.blob.decode("utf-8", "ignore")
        for part in prs.part.package.iter_parts()
        if part.partname.endswith(".xml")
    )
    assert "{{" not in xml  # every placeholder filled, the content marker removed
    assert "Khách hàng mẫu" in xml and "Công ty ABC (mẫu)" in xml
    assert prs.slides[1].shapes.title.text == "Về ABC"


async def test_docx_body_goes_to_template_marker() -> None:
    run = await finished()
    kit = ExportKit.load(source="sample")
    kit.logo = None  # without a logo the {{company_logo}} marker falls back to the short name
    doc = Document(io.BytesIO(build_docx(run, kit)))
    texts = [p.text for p in doc.paragraphs]
    assert "{{SCOPEAI_BODY}}" not in texts and not any("{{" in t for t in texts)
    cover = texts.index(run.proposal.title)
    assert cover < texts.index("Về ABC") < texts.index("Phụ lục")
    assert "ABC" in doc.sections[0].header.paragraphs[0].text


async def test_workbook_fills_template_mapping() -> None:
    run = await finished()
    kit = ExportKit.load(source="sample")
    wb = load_workbook(io.BytesIO(build_workbook(run, kit)))
    assert wb.sheetnames[:3] == ["Cover", "WBS", "Pricing"] and "Báo giá" in wb.sheetnames
    wbs, pricing = wb["WBS"], wb["Pricing"]
    first = next(i for i in run.wbs.items if i.level == 2)  # one template row per level-2 task
    assert wbs["A6"].value == first.id and wbs["E6"].value == first.estimate_md
    assert wbs["D6"].value and wbs["F6"].value == 1  # work types, project week
    assert pricing["E6"].value == "=C6*D6" and pricing["E3"].value == "=SUM(E6:E500)"
    assert "{{" not in str(wb["Cover"]["A5"].value)


async def test_qa_sheet_template_roundtrip() -> None:
    run = await finished("Chúng tôi muốn dùng AI để tăng năng suất cho nhân viên văn phòng.")
    kit = ExportKit.load(source="sample")
    wb = load_workbook(io.BytesIO(build_qa_sheet(run, kit)))
    ws = wb["Q&A"]
    first = run.gaps.questions[0]
    assert ws["A6"].value == first.id and ws["D6"].value == first.question
    ws["F6"] = "Khoảng 50 người"
    buffer = io.BytesIO()
    wb.save(buffer)
    assert parse_qa_answers(buffer.getvalue(), kit.config.qa_sheet) == {first.id: "Khoảng 50 người"}


def test_upload_validation_and_active_source(kb) -> None:
    with pytest.raises(TemplateError):
        save_custom("slides", "deck.pptx", b"not a pptx")
    with pytest.raises(TemplateError):
        save_custom("slides", "deck.docx", b"x")
    with pytest.raises(TemplateError):
        set_active("slides", "custom")  # nothing uploaded yet
    plain = io.BytesIO()
    Presentation().save(plain)
    warnings = save_custom("slides", "Mau.PPTX", plain.getvalue())
    assert any("SCOPEAI_CONTENT" in w for w in warnings)  # still accepted, region computed
    assert load_templates_config().active["slides"] == "custom"
    assert ExportKit.load().templates["slides"] == template_path("slides", "custom")
    assert delete_custom("slides") and load_templates_config().active["slides"] == "sample"
    assert not delete_custom("slides")


async def test_plain_pptx_template_still_builds(kb) -> None:
    plain = io.BytesIO()
    Presentation().save(plain)  # 4:3, no marker, default layout names
    save_custom("slides", "plain.pptx", plain.getvalue())
    run = await finished()
    prs = Presentation(io.BytesIO(build_slides(run)))
    assert len(prs.slides) == 15 and prs.slide_width < 10_000_000  # kept the template's size
    with zipfile.ZipFile(io.BytesIO(build_slides(run))) as z:
        assert not any("{{" in z.read(n).decode("utf-8", "ignore") for n in z.namelist())


# ---------- company logo ----------
def test_logo_upload_validation(kb) -> None:
    for name, data, message in [
        ("logo.gif", image_bytes((200, 80)), "PNG hoặc JPG"),
        ("logo.png", b"not an image", "không phải ảnh"),
        ("logo.png", image_bytes((10, 10)), "quá nhỏ"),
        ("logo.png", b"x" * (2 * 1024 * 1024 + 1), "2 MB"),
    ]:
        with pytest.raises(TemplateError, match=message):
            save_logo(name, data)
    assert not logo_path().exists()
    info = save_logo("logo.png", image_bytes((3000, 600)))
    assert (info["width"], info["height"]) == (1200, 240)  # scaled down, aspect kept
    assert save_logo("logo.JPG", image_bytes((400, 100), "JPEG"))["width"] == 400
    with Image.open(logo_path()) as img:
        assert img.format == "PNG"  # always re-encoded
    assert ExportKit.load().logo == logo_path()
    assert delete_logo() and not delete_logo()
    assert ExportKit.load().logo is None


def _pictures(shapes) -> int:
    return sum(1 for shape in shapes if shape.shape_type == MSO_SHAPE_TYPE.PICTURE)


async def test_logo_in_builtin_and_sample_exports(kb) -> None:
    save_logo("logo.png", image_bytes((400, 120)))
    run = await finished()

    prs = Presentation(io.BytesIO(build_slides(run, kit=ExportKit.load(source="builtin"))))
    assert _pictures(prs.slides[0].shapes) == 1 and _pictures(prs.slides[1].shapes) == 1

    prs = Presentation(io.BytesIO(build_slides(run, kit=ExportKit.load(source="sample"))))
    assert _pictures(prs.slide_master.shapes) == 1
    cover = next(lo for lo in prs.slide_layouts if lo.name == "Title Slide")
    assert _pictures(cover.shapes) == 1
    xml = "".join(
        part.blob.decode("utf-8", "ignore")
        for part in prs.part.package.iter_parts()
        if part.partname.endswith(".xml")
    )
    assert "{{" not in xml

    for source in ("builtin", "sample"):
        doc = Document(io.BytesIO(build_docx(run, ExportKit.load(source=source))))
        header = doc.sections[0].header
        assert "pic:pic" in header._element.xml and "{{" not in header.paragraphs[0].text
