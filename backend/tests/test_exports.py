import io
import zipfile

from docx import Document
from openpyxl import load_workbook
from pptx import Presentation
from pptx.util import Inches

from app.agents.pipeline import run_pipeline
from app.agents.pricing import compute_quotation, load_rate_card
from app.exports.common import parse_mermaid
from app.exports.document import build_docx
from app.exports.i18n import LABELS
from app.exports.package import build_package
from app.exports.slides import build_slides, deck_texts
from app.exports.workbook import build_workbook
from app.ingest import build_attachment
from app.llm.mock import MockLLM
from tests.conftest import MemoryRepo, make_run
from tests.samples import code_zip, data_csv, requirements_xlsx


async def finished_run(with_attachments: bool = True):
    run = make_run()
    if with_attachments:
        run.attachments = [
            build_attachment("a1", "req.xlsx", requirements_xlsx()),
            build_attachment("a2", "data.csv", data_csv()),
            build_attachment("a3", "code.zip", code_zip()),
        ]
    async for _ in run_pipeline(MockLLM(), run, MemoryRepo()):
        pass
    return run


def test_parse_mermaid_nodes_edges_and_chains() -> None:
    graph = parse_mermaid(
        "flowchart LR\n  U[Người dùng] --> UI[Web] -->|REST| API(API)\n  API --> DB[(DB)]"
    )
    assert graph.nodes["U"] == "Người dùng" and graph.nodes["API"] == "API"
    assert (
        ("U", "UI") in graph.edges and ("UI", "API") in graph.edges and ("API", "DB") in graph.edges
    )


async def test_slides_contain_all_sections() -> None:
    run = await finished_run()
    prs = Presentation(io.BytesIO(build_slides(run)))
    text = " ".join(
        shape.text_frame.text
        for slide in prs.slides
        for shape in slide.shapes
        if shape.has_text_frame
    )
    # 13 analysis slides + team + "Về ABC" + case studies + standard terms
    assert len(prs.slides) == 17
    for heading in [
        "Về ABC",
        "Đội ngũ dự án",
        "Phạm vi KHÔNG bao gồm",
        "Dự án tương tự đã triển khai",
        "Phương pháp, chất lượng & điều khoản",
        "Chi phí dự kiến",
        "Kiến trúc hệ thống",
        "Timeline triển khai",
        "Đáp ứng yêu cầu",
        "WBS",
        "Tài liệu & dữ liệu",
    ]:
        assert heading in text


async def test_slides_skip_optional_sections_without_attachments() -> None:
    run = await finished_run(with_attachments=False)
    prs = Presentation(io.BytesIO(build_slides(run)))
    assert len(prs.slides) == 15


async def test_team_slide_comes_from_wbs_and_quotation() -> None:
    run = await finished_run()
    prs = Presentation(io.BytesIO(build_slides(run)))
    team = next(
        s
        for s in prs.slides
        if any(sh.has_text_frame and sh.text_frame.text == "Đội ngũ dự án" for sh in s.shapes)
    )
    table = next(sh.table for sh in team.shapes if sh.has_table)
    cells = [cell.text for row in table.rows for cell in row.cells]
    labels = LABELS["vi"]["type_labels"]
    assert {labels[t.value] for total in run.wbs.totals for t in total.by_type} <= set(cells)
    assert "FPT" not in " ".join(cells)
    run.intake.language = "ja"  # a Japanese client adds the BrSE overhead row
    run.quotation = compute_quotation(run, load_rate_card())
    prs = Presentation(io.BytesIO(build_slides(run)))
    texts = " ".join(sh.text_frame.text for s in prs.slides for sh in s.shapes if sh.has_text_frame)
    tables = " ".join(
        c.text for s in prs.slides for sh in s.shapes if sh.has_table
        for row in sh.table.rows for c in row.cells
    )  # fmt: skip
    brse = next(ln.role_label for ln in run.quotation.lines if ln.kind == "overhead")
    assert brse in tables and texts


async def test_out_of_scope_is_translated_and_above_footer() -> None:
    run = await finished_run()
    assert set(run.architecture.out_of_scope) <= set(deck_texts(run))
    prs = Presentation(io.BytesIO(build_slides(run)))
    box = next(
        sh
        for s in prs.slides
        for sh in s.shapes
        if sh.has_text_frame and sh.text_frame.text.startswith("Phạm vi KHÔNG bao gồm")
    )
    assert box.top + box.height <= Inches(7.0)  # footer starts at 7.0 in


async def test_workbook_has_formulas_and_matrix() -> None:
    run = await finished_run()
    wb = load_workbook(io.BytesIO(build_workbook(run)))
    assert wb.sheetnames == [
        "Tổng quan",
        "Effort",
        "Báo giá",
        "WBS",
        "Timeline",
        "Đáp ứng yêu cầu",
        "Rủi ro",
        "Dữ liệu mẫu",
    ]
    quote = wb["Báo giá"]
    quote_formulas = [
        c.value
        for row in quote.iter_rows()
        for c in row
        if isinstance(c.value, str) and c.value.startswith("=")
    ]
    assert any(f.startswith("=SUM(E") for f in quote_formulas)
    assert any("*C" in f and f.startswith("=ROUND(") for f in quote_formulas)  # milestones
    effort = wb["Effort"]
    formulas = [
        c.value
        for row in effort.iter_rows()
        for c in row
        if isinstance(c.value, str) and c.value.startswith("=")
    ]
    assert any(f.startswith("=ROUND(") for f in formulas)
    matrix = wb["Đáp ứng yêu cầu"]
    assert matrix["A2"].value == "FR-01" and matrix["E2"].value


async def test_docx_has_proposal_and_appendix() -> None:
    run = await finished_run()
    doc = Document(io.BytesIO(build_docx(run)))
    text = "\n".join(p.text for p in doc.paragraphs)
    assert (
        "Giả định" in text
        and "Phụ lục" in text
        and "C. WBS" in text
        and "D. Bảng đáp ứng yêu cầu" in text
    )
    assert len(doc.tables) >= 4


async def test_package_zip_lists_all_deliverables() -> None:
    run = await finished_run()
    names = zipfile.ZipFile(io.BytesIO(build_package(run))).namelist()
    assert set(names) == {
        "README.txt",
        "slides.pptx",
        "proposal.docx",
        "workbook.xlsx",
        "proposal.md",
        "architecture.mmd",
        "qa_sheet.xlsx",
        "bidding.xlsx",
    }
