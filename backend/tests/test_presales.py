"""Pricing, deal stages, Q&A sheet, similarity, settings store and localized slides."""

import io
import shutil
from datetime import UTC, datetime

import pytest
from pptx import Presentation

from app.agents.pipeline import run_pipeline
from app.agents.pricing import compute_quotation, load_rate_card, match_role
from app.deal import sync_deal_stage
from app.documents import DocumentError
from app.exports.qa_sheet import build_qa_sheet, parse_qa_answers
from app.exports.slides import build_slides, deck_texts
from app.exports.translation import ensure_deck_translations
from app.knowledge import loader
from app.llm.mock import MockLLM
from app.schemas.deal import DealStage, PricingApproval
from app.schemas.quotation import Quotation
from app.schemas.run import RunStatus
from app.settings_store import (
    delete_reference_project,
    read_settings,
    save_estimation_template,
    save_rate_card,
    save_reference_project,
)
from app.similar import find_similar, similarity
from tests.conftest import MemoryRepo, make_run


async def finished(text: str = "Yêu cầu giả lập đủ dài để vượt qua kiểm tra độ dài."):
    run = make_run(text)
    async for _ in run_pipeline(MockLLM(), run, MemoryRepo()):
        pass
    return run


# ---------- pricing ----------
def test_match_role_keywords_and_default() -> None:
    card = load_rate_card()
    assert match_role("AI engineer", card).key == "ai_engineer"
    assert match_role("Senior Backend Developer", card).key == "backend"
    assert match_role("Data engineer", card).key == "data_engineer"
    assert match_role("QA", card).key == "qa"
    assert match_role("Chuyên gia tư vấn", card).key == card.default_role
    assert match_role("Maintainer", card).key == card.default_role  # "ai" must be a whole word


async def test_quotation_is_consistent() -> None:
    run = await finished()
    q = run.quotation
    assert q is not None and q.currency == "VND"
    assert q.subtotal == sum(p.amount for p in q.phases) == sum(line.amount for line in q.lines)
    wbs_days = sum(t.total_md for t in run.wbs.totals)
    assert sum(line.person_days for line in q.lines if line.kind == "wbs") == wbs_days
    assert q.wbs_person_days == wbs_days
    # PM is planned in the WBS (mandatory task), so the 10% PM overhead is not added on top;
    # no BrSE either for a Vietnamese client
    assert any(line.role_key == "pm" and line.kind == "wbs" for line in q.lines)
    assert q.overhead_person_days == 0 and not [ln for ln in q.lines if ln.kind == "overhead"]
    assert q.contract_model == "fixed_price" and q.onsite_ratio == 0
    assert q.total == q.subtotal + q.contingency
    assert q.total_min <= q.total <= q.total_max
    assert q.contingency_pct == 15  # base 10 + one risk with severity 4
    assert sum(m.percent for m in q.milestones) == 100
    assert q.monthly_run_cost == load_rate_card().run_cost_monthly["rag"]  # rag, cloud


async def test_quotation_currency_and_override() -> None:
    run = await finished()
    card = load_rate_card()
    vnd = compute_quotation(run, card, "VND", 0)
    jpy = compute_quotation(run, card, "JPY", 0)
    assert vnd and jpy and vnd.contingency == 0 and vnd.total == vnd.subtotal
    assert jpy.currency == "JPY" and abs(jpy.subtotal * 170 - vnd.subtotal) / vnd.subtotal < 0.01
    assert all(line.amount % 100 == 0 for line in jpy.lines)


async def test_quotation_contract_models_and_onsite() -> None:
    run = await finished()
    card = load_rate_card()
    fixed = compute_quotation(run, card, "VND")
    tm = compute_quotation(run, card, "VND", contract_model="time_material")
    odc = compute_quotation(run, card, "VND", contract_model="odc")
    onsite = compute_quotation(run, card, "VND", onsite_ratio=50)
    assert fixed and tm and odc and onsite
    assert tm.contingency == 0 and tm.total == tm.subtotal < fixed.total
    assert (
        len(tm.milestones) == tm.months and abs(sum(m.percent for m in tm.milestones) - 100) < 0.1
    )
    assert odc.odc_team and odc.odc_monthly_cost == sum(m.monthly_cost for m in odc.odc_team)
    assert odc.total == odc.odc_monthly_cost * odc.months
    assert onsite.subtotal > fixed.subtotal * 1.5  # onsite day rate = 3x offshore


async def test_quotation_adds_brse_for_japanese_client() -> None:
    run = await finished()
    run.intake.language = "ja"
    q = compute_quotation(run, load_rate_card())
    assert q and q.currency == "JPY"
    assert {line.role_key for line in q.lines if line.kind == "overhead"} == {"bridge_se"}


async def test_no_quotation_without_wbs() -> None:
    run = make_run()
    assert compute_quotation(run, load_rate_card()) is None


# ---------- deal stage ----------
def test_deal_stage_follows_status_until_manual() -> None:
    run = make_run()
    for status, stage in [
        (RunStatus.WAITING_CLARIFICATION, DealStage.CLARIFYING),
        (RunStatus.DONE, DealStage.REVIEWING),
        (RunStatus.APPROVED, DealStage.READY),
    ]:
        run.status = status
        sync_deal_stage(run)
        assert run.deal_stage == stage
    run.deal_stage = DealStage.REVIEWING
    run.quotation = Quotation.model_construct()  # a quote needs the pricing sign-off first
    sync_deal_stage(run)
    assert run.deal_stage == DealStage.REVIEWING
    run.pricing_approval = PricingApproval(approved=True, at=datetime.now(UTC))
    sync_deal_stage(run)
    assert run.deal_stage == DealStage.READY
    run.deal_stage = DealStage.WON
    run.status = RunStatus.DONE
    sync_deal_stage(run)
    assert run.deal_stage == DealStage.WON


# ---------- Q&A sheet ----------
@pytest.mark.parametrize(
    ("lang", "answer_header"),
    [("vi", "Câu trả lời của Quý khách"), ("ja", "ご回答"), ("en", "Your answer")],
)
async def test_qa_sheet_roundtrip(lang: str, answer_header: str) -> None:
    from openpyxl import load_workbook

    run = await finished()
    run.intake.language = lang
    data = build_qa_sheet(run)
    wb = load_workbook(io.BytesIO(data))
    ws = wb.active
    headers = [c.value for c in ws[4]]
    assert headers[0] == "ID" and headers[-1] == answer_header
    ws.cell(row=5, column=len(headers), value="Cập nhật hằng tuần")
    buffer = io.BytesIO()
    wb.save(buffer)
    assert parse_qa_answers(buffer.getvalue()) == {"q1": "Cập nhật hằng tuần"}


def test_qa_parse_rejects_unrelated_file() -> None:
    from tests.samples import requirements_xlsx

    with pytest.raises(DocumentError, match="Không tìm thấy bảng câu hỏi"):
        parse_qa_answers(requirements_xlsx())


# ---------- similarity ----------
async def test_similar_runs_ranked() -> None:
    a = await finished()
    b = await finished()
    c = await finished("Hàng tháng kế toán nhận 40 file Excel hóa đơn, muốn dùng AI tính tổng.")
    b.id, c.id = "bbbb0000", "cccc0000"
    assert similarity(a, b) > similarity(a, c)
    result = find_similar(a, [a, b, c])
    assert [r.id for r, _ in result][0] == "bbbb0000"
    assert all(r.id != a.id for r, _ in result)


# ---------- settings store ----------
@pytest.fixture
def kb_copy(tmp_path, monkeypatch):
    target = tmp_path / "kb"
    shutil.copytree(loader.KB_DIR, target)
    monkeypatch.setattr(loader, "KB_DIR", target)
    loader.reload()
    yield target
    loader.reload()


def test_settings_roundtrip_and_validation(kb_copy) -> None:
    data = read_settings()
    card = data["rate_card"]
    card["roles"][0]["day_rate"] = 5_000_000
    save_rate_card(card)
    assert read_settings()["rate_card"]["roles"][0]["day_rate"] == 5_000_000

    bad = {**card, "milestones": [{"name": "A", "percent": 50}]}
    with pytest.raises(ValueError, match="100"):
        save_rate_card(bad)

    template = data["estimation_template"]
    template["base"]["rag"]["mvp"] = [40, 30]
    with pytest.raises(ValueError, match="min ≤ max"):
        save_estimation_template(template)

    save_reference_project(
        "rp_test_case", "# rp_test_case\\n- Pattern: rag, dự án giả lập để kiểm thử"
    )
    assert any(r["name"] == "rp_test_case" for r in read_settings()["reference_projects"])
    with pytest.raises(ValueError):
        save_reference_project("Bad Name!", "x" * 30)
    assert delete_reference_project("rp_test_case") and not delete_reference_project("rp_test_case")


# ---------- localized slides ----------
async def test_japanese_deck_uses_labels_and_translator() -> None:
    run = await finished()
    seen: list[str] = []

    def tr(text: str) -> str:
        seen.append(text)
        return f"[JA]{text}"

    prs = Presentation(io.BytesIO(build_slides(run, "ja", tr)))
    text = " ".join(
        s.text_frame.text for slide in prs.slides for s in slide.shapes if s.has_text_frame
    )
    assert "概算お見積り" in text and "システム構成" in text and "[JA]" in text
    assert run.pattern.rationale in seen


async def test_translation_cache_with_mock_identity() -> None:
    run = await finished()
    llm = MockLLM()
    tr = await ensure_deck_translations(run, "en", llm)
    assert tr(run.pattern.rationale) == run.pattern.rationale  # mock cannot translate
    assert set(deck_texts(run)) <= set(run.deck_translations["en"])
    calls = len(llm.calls)
    await ensure_deck_translations(run, "en", llm)
    assert len(llm.calls) == calls  # cached
